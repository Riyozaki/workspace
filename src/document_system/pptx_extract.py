from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .errors import UnsupportedDocumentError
from .security import detect_format

PLACEHOLDER_RE = re.compile(r"\{\{[^{}\n]{1,120}\}\}|<<[^<>\n]{1,120}>>|\b(?:TODO|TBD)\b", re.IGNORECASE)


def _inches(value: int) -> float:
    return round(value / 914400, 4)


def _shape_text(shape: Any) -> str:
    if getattr(shape, "has_text_frame", False):
        return shape.text
    if getattr(shape, "has_table", False):
        return "\n".join(
            "\t".join(cell.text for cell in row.cells)
            for row in shape.table.rows
        )
    return ""


def extract_pptx(path: str | Path) -> dict[str, Any]:
    source = Path(path)
    if detect_format(source) != "pptx":
        raise UnsupportedDocumentError(f"Not a PPTX package: {source}")
    presentation = Presentation(source)
    slides: list[dict[str, Any]] = []
    all_placeholders: list[dict[str, Any]] = []
    totals = {
        "slides": len(presentation.slides),
        "shapes": 0,
        "text_shapes": 0,
        "pictures": 0,
        "tables": 0,
        "charts": 0,
        "notes_slides": 0,
    }

    for slide_number, slide in enumerate(presentation.slides, start=1):
        shapes: list[dict[str, Any]] = []
        title = None
        slide_texts: list[str] = []
        for shape_index, shape in enumerate(slide.shapes, start=1):
            text = _shape_text(shape)
            if text:
                slide_texts.append(text)
            if shape.name == "ds-role:title" and text:
                title = text
            shape_type = str(shape.shape_type)
            item: dict[str, Any] = {
                "index": shape_index,
                "name": shape.name,
                "type": shape_type,
                "x": _inches(shape.left),
                "y": _inches(shape.top),
                "width": _inches(shape.width),
                "height": _inches(shape.height),
                "text": text or None,
            }
            if getattr(shape, "has_text_frame", False):
                runs = []
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        runs.append(
                            {
                                "text": run.text,
                                "font": run.font.name,
                                "size_pt": round(run.font.size.pt, 2) if run.font.size else None,
                                "bold": run.font.bold,
                                "italic": run.font.italic,
                            }
                        )
                item["runs"] = runs
                item["paragraphs"] = len(shape.text_frame.paragraphs)
                totals["text_shapes"] += 1
            if getattr(shape, "has_table", False):
                item["rows"] = len(shape.table.rows)
                item["columns"] = len(shape.table.columns)
                totals["tables"] += 1
            if getattr(shape, "has_chart", False):
                item["chart_type"] = str(shape.chart.chart_type)
                item["series"] = len(shape.chart.series)
                totals["charts"] += 1
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                item["image_sha256"] = hashlib.sha256(shape.image.blob).hexdigest()
                item["image_format"] = shape.image.ext
                totals["pictures"] += 1
            for match in PLACEHOLDER_RE.finditer(text):
                placeholder = {
                    "slide": slide_number,
                    "shape": shape_index,
                    "value": match.group(0),
                }
                all_placeholders.append(placeholder)
            shapes.append(item)
            totals["shapes"] += 1

        if title is None:
            title = next((text for text in slide_texts if text.strip()), None)
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                totals["notes_slides"] += 1
        slides.append(
            {
                "number": slide_number,
                "title": title,
                "layout": slide.slide_layout.name,
                "layout_part": str(slide.slide_layout.part.partname),
                "master_part": str(slide.slide_layout.slide_master.part.partname),
                "text": "\n".join(slide_texts),
                "shape_count": len(shapes),
                "shapes": shapes,
                "notes": notes or None,
            }
        )

    masters = []
    for master_index, master in enumerate(presentation.slide_masters, start=1):
        masters.append(
            {
                "index": master_index,
                "part": str(master.part.partname),
                "layouts": [
                    {
                        "index": layout_index,
                        "name": layout.name,
                        "part": str(layout.part.partname),
                    }
                    for layout_index, layout in enumerate(master.slide_layouts, start=1)
                ],
            }
        )

    return {
        "path": str(source),
        "slide_width": _inches(presentation.slide_width),
        "slide_height": _inches(presentation.slide_height),
        "slides": slides,
        "masters": masters,
        "totals": totals,
        "placeholders": all_placeholders,
        "properties": {
            "title": presentation.core_properties.title,
            "subject": presentation.core_properties.subject,
            "author": presentation.core_properties.author,
            "keywords": presentation.core_properties.keywords,
            "category": presentation.core_properties.category,
            "comments": presentation.core_properties.comments,
        },
    }
