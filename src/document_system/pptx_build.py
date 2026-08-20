from __future__ import annotations

import importlib.resources
import json
import math
import zipfile
from pathlib import Path
from typing import Any

import jsonschema
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from .util import read_json

DEFAULT_THEME = {
    "primary": "203F54",
    "secondary": "3E6B78",
    "accent": "C77932",
    "background": "F7F9FA",
    "light": "E9F0F3",
    "text": "23313A",
    "muted": "667985",
    "positive": "25734C",
    "negative": "B04444",
    "font": "Arial",
    "heading_font": "Arial",
}

CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "area": XL_CHART_TYPE.AREA,
}

LEGEND_POSITIONS = {
    "top": XL_LEGEND_POSITION.TOP,
    "bottom": XL_LEGEND_POSITION.BOTTOM,
    "left": XL_LEGEND_POSITION.LEFT,
    "right": XL_LEGEND_POSITION.RIGHT,
}


def _rgb(value: str) -> RGBColor:
    value = value.strip().lstrip("#").upper()
    if len(value) != 6 or any(char not in "0123456789ABCDEF" for char in value):
        raise ValueError(f"Expected six-digit RGB color, got {value!r}")
    return RGBColor.from_string(value)


def validate_pptx_spec(spec: dict[str, Any]) -> None:
    schema_path = importlib.resources.files("document_system").joinpath("schemas/pptx-spec.schema.json")
    schema = json.loads(schema_path.read_text(encoding="utf-8"))
    jsonschema.Draft202012Validator(schema).validate(spec)
    for index, slide in enumerate(spec["slides"], start=1):
        layout = slide["layout"]
        required = {
            "title": ("title",),
            "section": ("title",),
            "bullets": ("title", "bullets"),
            "two-column": ("title", "left", "right"),
            "metrics": ("title", "metrics"),
            "table": ("title", "headers", "rows"),
            "chart": ("title", "chart_type", "categories", "series"),
            "image": ("title", "image"),
            "quote": ("quote",),
            "timeline": ("title", "events"),
        }[layout]
        missing = [key for key in required if key not in slide]
        if missing:
            raise ValueError(f"Slide {index} ({layout}) is missing: {', '.join(missing)}")
        if layout == "table":
            columns = len(slide["headers"])
            if not columns:
                raise ValueError(f"Slide {index}: table must have headers")
            if any(len(row) > columns for row in slide["rows"]):
                raise ValueError(f"Slide {index}: table row has more values than headers")
            widths = slide.get("widths_pct")
            if widths and (len(widths) != columns or not math.isclose(sum(widths), 100, abs_tol=1.0)):
                raise ValueError(f"Slide {index}: widths_pct must match columns and sum to 100")
        if layout == "chart":
            count = len(slide["categories"])
            if any(len(series["values"]) != count for series in slide["series"]):
                raise ValueError(f"Slide {index}: every chart series must match category count")


def _set_run(run: Any, theme: dict[str, str], language: str, *, size: float, color: str, bold: bool = False, italic: bool = False, font: str | None = None) -> None:
    run.font.name = font or theme["font"]
    run.font.size = Pt(size)
    run.font.color.rgb = _rgb(color)
    run.font.bold = bold
    run.font.italic = italic
    run._r.get_or_add_rPr().set("lang", language)


def _add_text(
    slide: Any,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    theme: dict[str, str],
    language: str,
    *,
    size: float = 18,
    color: str | None = None,
    bold: bool = False,
    italic: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
    margin: float = 0.04,
    name: str = "ds-role:text",
) -> Any:
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    _set_run(run, theme, language, size=size, color=color or theme["text"], bold=bold, italic=italic)
    return shape


def _add_bullets(slide: Any, items: list[Any], x: float, y: float, w: float, h: float, theme: dict[str, str], language: str, *, font_size: float = 20) -> Any:
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = "ds-role:bullets"
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.05)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.03)
    first = True
    for item in items:
        config = item if isinstance(item, dict) else {"text": item, "level": 0}
        paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        level = int(config.get("level", 0))
        paragraph.level = level
        paragraph.space_after = Pt(9 if level == 0 else 5)
        paragraph.line_spacing = 1.08
        ppr = paragraph._p.get_or_add_pPr()
        for tag in ("{http://schemas.openxmlformats.org/drawingml/2006/main}buNone", "{http://schemas.openxmlformats.org/drawingml/2006/main}buChar"):
            for node in list(ppr.findall(tag)):
                ppr.remove(node)
        bullet = ppr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar", {"char": "•"})
        ppr.append(bullet)
        run = paragraph.add_run()
        run.text = config["text"]
        _set_run(
            run,
            theme,
            language,
            size=font_size - level * 2,
            color=theme["text"],
            bold=bool(config.get("emphasis")),
        )
    return shape


def _add_rect(slide: Any, x: float, y: float, w: float, h: float, fill: str, *, line: str | None = None, radius: bool = False, name: str = "ds-role:decor") -> Any:
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill)
    if line:
        shape.line.color.rgb = _rgb(line)
    else:
        shape.line.fill.background()
    return shape


def _set_background(slide: Any, color: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(color)


def _add_standard_title(slide: Any, title: str, theme: dict[str, str], language: str, number: int, footer: str | None, source: str | None) -> None:
    _add_text(slide, title, 0.65, 0.38, 11.7, 0.65, theme, language, size=25, color=theme["primary"], bold=True, name="ds-role:title")
    _add_rect(slide, 0.67, 1.08, 1.0, 0.06, theme["accent"], name="ds-role:accent")
    footer_text = footer or ""
    if footer_text:
        _add_text(slide, footer_text, 0.68, 7.12, 9.5, 0.22, theme, language, size=8.5, color=theme["muted"], name="ds-role:footer")
    if source:
        _add_text(slide, source, 0.68, 6.82, 11.2, 0.22, theme, language, size=8, color=theme["muted"], name="ds-role:source")
    _add_text(slide, str(number), 12.18, 7.08, 0.45, 0.24, theme, language, size=9, color=theme["muted"], align=PP_ALIGN.RIGHT, name="ds-role:slide-number")


def _add_picture(slide: Any, path: Path, x: float, y: float, w: float, h: float, *, fit: str = "cover") -> Any:
    with Image.open(path) as image:
        iw, ih = image.size
    image_ratio = iw / ih
    box_ratio = w / h
    if fit == "contain":
        if image_ratio >= box_ratio:
            actual_w, actual_h = w, w / image_ratio
        else:
            actual_w, actual_h = h * image_ratio, h
        picture = slide.shapes.add_picture(
            str(path), Inches(x + (w - actual_w) / 2), Inches(y + (h - actual_h) / 2), Inches(actual_w), Inches(actual_h)
        )
    else:
        picture = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
        if image_ratio > box_ratio:
            visible = box_ratio / image_ratio
            picture.crop_left = picture.crop_right = (1 - visible) / 2
        elif image_ratio < box_ratio:
            visible = image_ratio / box_ratio
            picture.crop_top = picture.crop_bottom = (1 - visible) / 2
    picture.name = "ds-role:image"
    return picture


def _render_slide(slide: Any, config: dict[str, Any], theme: dict[str, str], language: str, number: int, global_footer: str, base_dir: Path) -> None:
    layout = config["layout"]
    background = config.get("background", theme["background"])
    _set_background(slide, background)

    if layout == "title":
        _add_rect(slide, 0, 0, 4.1, 7.5, theme["primary"], name="ds-role:panel")
        _add_rect(slide, 4.1, 0, 0.13, 7.5, theme["accent"], name="ds-role:accent")
        if config.get("kicker"):
            _add_text(slide, config["kicker"].upper(), 4.75, 1.15, 7.6, 0.35, theme, language, size=11, color=theme["accent"], bold=True)
        _add_text(slide, config["title"], 4.7, 1.65, 7.7, 2.1, theme, language, size=32, color=theme["primary"], bold=True, valign=MSO_ANCHOR.MIDDLE, name="ds-role:title")
        if config.get("subtitle"):
            _add_text(slide, config["subtitle"], 4.75, 4.05, 7.1, 1.0, theme, language, size=17, color=theme["muted"])
        _add_text(slide, config.get("footer", global_footer), 4.75, 6.72, 6.8, 0.3, theme, language, size=9, color=theme["muted"], name="ds-role:footer")
    elif layout == "section":
        _set_background(slide, theme["primary"])
        _add_rect(slide, 0.72, 1.28, 0.14, 4.8, theme["accent"], name="ds-role:accent")
        if config.get("section_number"):
            _add_text(slide, config["section_number"], 1.18, 1.35, 1.4, 0.7, theme, language, size=32, color=theme["accent"], bold=True)
        _add_text(slide, config["title"], 1.18, 2.2, 10.9, 1.7, theme, language, size=34, color="FFFFFF", bold=True, valign=MSO_ANCHOR.MIDDLE, name="ds-role:title")
        if config.get("subtitle"):
            _add_text(slide, config["subtitle"], 1.2, 4.15, 9.7, 0.9, theme, language, size=17, color=theme["light"])
    else:
        _add_standard_title(slide, config.get("title", ""), theme, language, number, config.get("footer", global_footer), config.get("source"))
        if layout == "bullets":
            _add_bullets(slide, config["bullets"], 0.85, 1.42, 11.55, 5.1, theme, language, font_size=21)
        elif layout == "two-column":
            _add_rect(slide, 0.72, 1.42, 5.85, 4.98, "FFFFFF", line=theme["light"], radius=True, name="ds-role:card")
            _add_rect(slide, 6.76, 1.42, 5.85, 4.98, "FFFFFF", line=theme["light"], radius=True, name="ds-role:card")
            if config.get("left_title"):
                _add_text(slide, config["left_title"], 1.0, 1.7, 5.3, 0.45, theme, language, size=17, color=theme["secondary"], bold=True)
            if config.get("right_title"):
                _add_text(slide, config["right_title"], 7.03, 1.7, 5.3, 0.45, theme, language, size=17, color=theme["secondary"], bold=True)
            _add_bullets(slide, config["left"], 1.0, 2.25, 5.25, 3.85, theme, language, font_size=17)
            _add_bullets(slide, config["right"], 7.03, 2.25, 5.25, 3.85, theme, language, font_size=17)
        elif layout == "metrics":
            metrics = config["metrics"]
            gap = 0.22
            card_width = (11.8 - gap * (len(metrics) - 1)) / len(metrics)
            for index, metric in enumerate(metrics):
                x = 0.75 + index * (card_width + gap)
                _add_rect(slide, x, 1.75, card_width, 3.9, "FFFFFF", line=theme["light"], radius=True, name="ds-role:metric-card")
                status = metric.get("status", "neutral")
                color = {"neutral": theme["secondary"], "positive": theme["positive"], "negative": theme["negative"], "accent": theme["accent"]}[status]
                _add_rect(slide, x, 1.75, card_width, 0.09, color, name="ds-role:accent")
                _add_text(slide, metric["value"], x + 0.24, 2.15, card_width - 0.48, 1.0, theme, language, size=28, color=color, bold=True, valign=MSO_ANCHOR.MIDDLE)
                _add_text(slide, metric["label"], x + 0.24, 3.28, card_width - 0.48, 0.75, theme, language, size=14, color=theme["text"], bold=True)
                if metric.get("detail"):
                    _add_text(slide, metric["detail"], x + 0.24, 4.22, card_width - 0.48, 0.8, theme, language, size=11.5, color=theme["muted"])
        elif layout == "table":
            headers = config["headers"]
            rows = config["rows"]
            shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(0.72), Inches(1.5), Inches(11.9), Inches(4.95))
            shape.name = "ds-role:table"
            table = shape.table
            widths = config.get("widths_pct") or [100 / len(headers)] * len(headers)
            for index, width in enumerate(widths):
                table.columns[index].width = Inches(11.9 * width / 100)
            for column, value in enumerate(headers):
                cell = table.cell(0, column)
                cell.text = str(value)
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(theme["primary"])
            for row_index, values in enumerate(rows, start=1):
                for column in range(len(headers)):
                    cell = table.cell(row_index, column)
                    cell.text = "" if column >= len(values) or values[column] is None else str(values[column])
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = _rgb("FFFFFF" if row_index % 2 else theme["light"])
            for row_index in range(len(rows) + 1):
                for column in range(len(headers)):
                    cell = table.cell(row_index, column)
                    cell.margin_left = cell.margin_right = Inches(0.08)
                    cell.margin_top = cell.margin_bottom = Inches(0.05)
                    frame = cell.text_frame
                    frame.word_wrap = True
                    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                    for paragraph in frame.paragraphs:
                        for run in paragraph.runs:
                            _set_run(run, theme, language, size=11 if row_index else 11.5, color="FFFFFF" if row_index == 0 else theme["text"], bold=row_index == 0)
        elif layout == "chart":
            chart_data = ChartData()
            chart_data.categories = config["categories"]
            for series in config["series"]:
                chart_data.add_series(series["name"], series["values"])
            chart = slide.shapes.add_chart(
                CHART_TYPES[config["chart_type"]],
                Inches(0.9), Inches(1.45), Inches(11.5), Inches(5.15), chart_data
            ).chart
            chart.has_title = False
            chart.has_legend = config.get("legend", "bottom") != "none"
            if chart.has_legend:
                chart.legend.position = LEGEND_POSITIONS[config.get("legend", "bottom")]
                chart.legend.include_in_layout = False
                chart.legend.font.name = theme["font"]
                chart.legend.font.size = Pt(10)
            if hasattr(chart, "value_axis") and config.get("y_title"):
                chart.value_axis.has_title = True
                chart.value_axis.axis_title.text_frame.text = config["y_title"]
            if config.get("data_labels"):
                plot = chart.plots[0]
                plot.has_data_labels = True
                plot.data_labels.show_value = True
                plot.data_labels.font.size = Pt(9)
            chart.chart_style = 10
        elif layout == "image":
            path = (base_dir / config["image"]).resolve()
            _add_picture(slide, path, 0.78, 1.4, 11.78, 5.22, fit=config.get("image_fit", "contain"))
            if config.get("caption"):
                _add_text(slide, config["caption"], 0.85, 6.56, 11.45, 0.24, theme, language, size=9, color=theme["muted"], align=PP_ALIGN.CENTER, name="ds-role:caption")
        elif layout == "quote":
            _add_text(slide, "“", 0.95, 1.05, 1.2, 1.1, theme, language, size=60, color=theme["accent"], bold=True)
            _add_text(slide, config["quote"], 1.55, 1.65, 10.3, 3.15, theme, language, size=25, color=theme["primary"], italic=True, valign=MSO_ANCHOR.MIDDLE, name="ds-role:quote")
            if config.get("attribution"):
                _add_text(slide, "— " + config["attribution"], 6.4, 5.1, 5.35, 0.5, theme, language, size=13, color=theme["muted"], align=PP_ALIGN.RIGHT)
        elif layout == "timeline":
            events = config["events"]
            x0, x1, y = 1.05, 12.25, 3.25
            _add_rect(slide, x0, y, x1 - x0, 0.05, theme["secondary"], name="ds-role:timeline-line")
            step = (x1 - x0) / max(1, len(events) - 1)
            for index, event in enumerate(events):
                x = x0 + index * step
                circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - 0.12), Inches(y - 0.11), Inches(0.26), Inches(0.26))
                circle.name = "ds-role:timeline-marker"
                circle.fill.solid()
                circle.fill.fore_color.rgb = _rgb(theme["accent"] if index in {0, len(events) - 1} else theme["secondary"])
                circle.line.fill.background()
                _add_text(slide, event["label"], x - 0.75, y - 0.72, 1.5, 0.35, theme, language, size=10, color=theme["accent"], bold=True, align=PP_ALIGN.CENTER)
                _add_text(slide, event["title"], x - 0.9, y + 0.42, 1.8, 0.65, theme, language, size=12.5, color=theme["text"], bold=True, align=PP_ALIGN.CENTER)
                if event.get("body"):
                    _add_text(slide, event["body"], x - 0.95, y + 1.15, 1.9, 1.0, theme, language, size=9.5, color=theme["muted"], align=PP_ALIGN.CENTER)

    if config.get("notes"):
        notes_frame = slide.notes_slide.notes_text_frame
        notes_frame.text = config["notes"]


def _normalize_chart_axis_ids(path: Path) -> None:
    """python-pptx may emit signed axis IDs; OOXML requires UInt32 values."""
    with zipfile.ZipFile(path) as source:
        infos = source.infolist()
        values = {info.filename: source.read(info.filename) for info in infos}
    changed = False
    for name in [
        value
        for value in values
        if value.startswith("ppt/charts/chart") and value.endswith(".xml")
    ]:
        root = etree.fromstring(values[name])
        part_changed = False
        for node in root.iter():
            if node.tag.rsplit("}", 1)[-1] not in {"axId", "crossAx"}:
                continue
            raw = node.get("val")
            if raw and (raw.startswith("-") or int(raw) > 2**31 - 1):
                # Chart axis IDs are semantically constrained to positive Int32 by
                # PowerPoint/Open XML SDK even though the XML type is unsigned.
                node.set("val", str(abs(int(raw)) & 0x7FFFFFFF))
                part_changed = True
        if part_changed:
            values[name] = etree.tostring(
                root, encoding="UTF-8", xml_declaration=True, standalone=True
            )
            changed = True
    if not changed:
        return
    temporary = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(temporary, "w") as destination:
        for info in infos:
            destination.writestr(info, values[info.filename])
    temporary.replace(path)


def build_pptx(spec: dict[str, Any], output_path: str | Path, *, base_dir: str | Path | None = None) -> Path:
    validate_pptx_spec(spec)
    destination = Path(output_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    base = Path(base_dir or ".").resolve()
    theme = {**DEFAULT_THEME, **spec.get("theme", {})}
    for key in ("primary", "secondary", "accent", "background", "light", "text", "muted", "positive", "negative"):
        _rgb(theme[key])
    language = spec.get("language", "ru-RU")

    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    properties = presentation.core_properties
    metadata = spec.get("metadata", {})
    for key in ("title", "subject", "author", "keywords", "category", "comments"):
        if key in metadata:
            setattr(properties, key, metadata[key])
    blank_layout = presentation.slide_layouts[6]
    for number, config in enumerate(spec["slides"], start=1):
        slide = presentation.slides.add_slide(blank_layout)
        _render_slide(slide, config, theme, language, number, spec.get("footer", ""), base)
    # Remove the initial default slide only if a nonstandard template supplied one.
    while len(presentation.slides) > len(spec["slides"]):
        slide_id = presentation.slides._sldIdLst[0]
        presentation.part.drop_rel(slide_id.rId)
        presentation.slides._sldIdLst.remove(slide_id)
    presentation.save(destination)
    _normalize_chart_axis_ids(destination)
    return destination


def build_pptx_from_file(spec_path: str | Path, output_path: str | Path) -> Path:
    source = Path(spec_path).resolve()
    return build_pptx(read_json(source), output_path, base_dir=source.parent)
