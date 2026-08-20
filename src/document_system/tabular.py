from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Any

from .xlsx_build import build_xlsx

INTEGER_RE = re.compile(r"^-?(?:0|[1-9][0-9]*)$")
FLOAT_RE = re.compile(r"^-?(?:0|[1-9][0-9]*)\.[0-9]+$")


def _unique_headers(values: list[str]) -> tuple[list[str], list[dict[str, str]]]:
    result = []
    seen: dict[str, int] = {}
    changes = []
    for index, raw in enumerate(values, start=1):
        base = raw.strip() or f"Column_{index}"
        key = base.casefold()
        count = seen.get(key, 0) + 1
        seen[key] = count
        value = base if count == 1 else f"{base}_{count}"
        result.append(value)
        if value != raw:
            changes.append({"original": raw, "normalized": value})
    return result, changes


def _convert(value: str) -> Any:
    stripped = value.strip()
    if stripped == "":
        return None
    if INTEGER_RE.match(stripped) and not (
        len(stripped.lstrip("-")) > 1 and stripped.lstrip("-").startswith("0")
    ):
        try:
            return int(stripped)
        except ValueError:
            pass
    if FLOAT_RE.match(stripped):
        try:
            return float(stripped)
        except ValueError:
            pass
    if stripped.casefold() in {"true", "false"}:
        return stripped.casefold() == "true"
    return stripped


def import_tabular_to_xlsx(
    input_path: str | Path,
    output_path: str | Path,
    *,
    delimiter: str | None = None,
    sheet_name: str = "Data",
    encoding: str = "utf-8-sig",
    max_rows: int = 1_000_000,
    max_columns: int = 16_384,
) -> dict[str, Any]:
    source = Path(input_path)
    with source.open("r", encoding=encoding, newline="") as sample_handle:
        sample = sample_handle.read(64_000)
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
        except csv.Error:
            delimiter = "\t" if source.suffix.lower() == ".tsv" else ","
    rows: list[tuple[int, list[str]]] = []
    blank_rows = 0
    with source.open("r", encoding=encoding, newline="") as handle:
        reader = csv.reader(handle, delimiter=delimiter)
        for number, row in enumerate(reader, start=1):
            if number > max_rows:
                raise ValueError(f"Tabular input exceeds {max_rows} row safety limit")
            if len(row) > max_columns:
                raise ValueError(f"Row {number} exceeds {max_columns} column safety limit")
            if not any(value.strip() for value in row):
                blank_rows += 1
                continue
            rows.append((number, row))
    if not rows:
        raise ValueError("Tabular input has no non-empty rows")
    width = max(len(row) for _number, row in rows)
    header_row = rows[0][1]
    headers, header_changes = _unique_headers(
        header_row + [""] * (width - len(header_row))
    )
    ragged_rows = []
    body = []
    for source_row, row in rows[1:]:
        if len(row) != width:
            ragged_rows.append({"row": source_row, "columns": len(row), "expected": width})
        normalized = row[:width] + [""] * max(0, width - len(row))
        body.append([_convert(value) for value in normalized])
    widths = []
    for column in range(width):
        values = [headers[column]] + ["" if row[column] is None else str(row[column]) for row in body[:5000]]
        widths.append(min(60, max(10, max(len(value) for value in values) + 2)))
    columns = []
    from openpyxl.utils import get_column_letter

    for index, column_width in enumerate(widths, start=1):
        columns.append({"column": get_column_letter(index), "width": column_width})
    spec = {
        "metadata": {
            "title": source.stem,
            "creator": "documentctl tabular import",
            "description": f"Imported from {source.name}",
        },
        "active_sheet": sheet_name,
        "sheets": [
            {
                "name": sheet_name,
                "freeze_panes": "A2",
                "show_gridlines": False,
                "columns": columns,
                "tables": [
                    {
                        "name": "ImportedData",
                        "start_cell": "A1",
                        "headers": headers,
                        "rows": body,
                        "style": "TableStyleMedium2",
                    }
                ],
                "print": {
                    "orientation": "landscape" if width > 6 else "portrait",
                    "paper_size": "A4",
                    "fit_to_width": 1,
                    "fit_to_height": 0,
                    "repeat_rows": "1:1",
                    "print_area": f"A1:{get_column_letter(width)}{len(body) + 1}",
                },
            }
        ],
    }
    build_xlsx(spec, output_path)
    return {
        "input": str(source),
        "output": str(Path(output_path)),
        "delimiter": "\\t" if delimiter == "\t" else delimiter,
        "rows": len(body),
        "columns": width,
        "blank_rows_removed": blank_rows,
        "ragged_rows": ragged_rows,
        "header_changes": header_changes,
        "headers": headers,
    }
