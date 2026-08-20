from __future__ import annotations

import re
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .errors import UnsupportedDocumentError
from .security import detect_format

FORMULA_ERROR_VALUES = {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#GETTING_DATA"}
PLACEHOLDER_RE = re.compile(r"\{\{[^{}\n]{1,120}\}\}|<<[^<>\n]{1,120}>>|\b(?:TODO|TBD)\b", re.IGNORECASE)


def extract_xlsx(path: str | Path, *, include_cells: bool = True, max_cells_per_sheet: int = 5_000) -> dict[str, Any]:
    source = Path(path)
    if detect_format(source) != "xlsx":
        raise UnsupportedDocumentError(f"Not an XLSX package: {source}")
    formula_book = load_workbook(source, data_only=False, read_only=False, keep_links=True)
    value_book = load_workbook(source, data_only=True, read_only=False, keep_links=True)
    sheets: list[dict[str, Any]] = []
    totals: Counter[str] = Counter()
    all_placeholders: list[dict[str, str]] = []
    all_errors: list[dict[str, Any]] = []
    all_missing_cached: list[dict[str, str]] = []

    try:
        for ws in formula_book.worksheets:
            value_ws = value_book[ws.title]
            cells: list[dict[str, Any]] = []
            formulas: list[dict[str, Any]] = []
            errors: list[dict[str, Any]] = []
            missing_cached: list[dict[str, str]] = []
            placeholders: list[dict[str, str]] = []
            nonempty = 0
            comments = 0
            hyperlinks = 0
            style_counts: Counter[str] = Counter()
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    nonempty += 1
                    if cell.comment is not None:
                        comments += 1
                    if cell.hyperlink is not None:
                        hyperlinks += 1
                    if cell.style_id:
                        style_counts[str(cell.style_id)] += 1
                    cached = value_ws[cell.coordinate].value
                    formula = cell.value if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")) else None
                    if formula:
                        item = {"cell": cell.coordinate, "formula": formula, "cached": cached}
                        formulas.append(item)
                        if cached is None:
                            missing = {"sheet": ws.title, "cell": cell.coordinate, "formula": str(formula)}
                            missing_cached.append(missing)
                            all_missing_cached.append(missing)
                        if isinstance(cached, str) and cached.upper() in FORMULA_ERROR_VALUES:
                            error = {"sheet": ws.title, **item, "error": cached.upper()}
                            errors.append(error)
                            all_errors.append(error)
                    value_for_check = cached if formula else cell.value
                    if isinstance(value_for_check, str):
                        for match in PLACEHOLDER_RE.finditer(value_for_check):
                            placeholder = {"sheet": ws.title, "cell": cell.coordinate, "value": match.group(0)}
                            placeholders.append(placeholder)
                            all_placeholders.append(placeholder)
                    if include_cells and len(cells) < max_cells_per_sheet:
                        try:
                            style_name = cell.style
                        except IndexError:
                            style_name = f"style_id:{cell.style_id}"
                        cells.append(
                            {
                                "cell": cell.coordinate,
                                "value": cell.value,
                                "cached": cached if formula else None,
                                "data_type": cell.data_type,
                                "number_format": cell.number_format,
                                "style": style_name,
                                "comment": cell.comment.text if cell.comment else None,
                                "hyperlink": cell.hyperlink.target if cell.hyperlink else None,
                            }
                        )
            sheet = {
                "name": ws.title,
                "state": ws.sheet_state,
                "dimensions": ws.calculate_dimension(),
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "nonempty_cells": nonempty,
                "formula_count": len(formulas),
                "formulas": formulas,
                "formula_errors": errors,
                "missing_cached_values": missing_cached,
                "comments": comments,
                "hyperlinks": hyperlinks,
                "merged_ranges": [str(value) for value in ws.merged_cells.ranges],
                "tables": [
                    {"name": table.name, "range": table.ref, "style": table.tableStyleInfo.name if table.tableStyleInfo else None}
                    for table in ws.tables.values()
                ],
                "charts": len(ws._charts),
                "images": len(ws._images),
                "data_validations": len(ws.data_validations.dataValidation),
                "conditional_format_ranges": len(ws.conditional_formatting),
                "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
                "auto_filter": ws.auto_filter.ref,
                "print_area": str(ws.print_area) if ws.print_area else None,
                "placeholders": placeholders,
                "style_counts": dict(style_counts),
                "cells": cells if include_cells else None,
                "cells_truncated": max(0, nonempty - len(cells)) if include_cells else None,
            }
            sheets.append(sheet)
            totals.update(
                {
                    "sheets": 1,
                    "visible_sheets": int(ws.sheet_state == "visible"),
                    "nonempty_cells": nonempty,
                    "formulas": len(formulas),
                    "formula_errors": len(errors),
                    "missing_cached_values": len(missing_cached),
                    "comments": comments,
                    "hyperlinks": hyperlinks,
                    "tables": len(ws.tables),
                    "charts": len(ws._charts),
                    "images": len(ws._images),
                }
            )

        defined_names = []
        for name, value in formula_book.defined_names.items():
            defined_names.append(
                {
                    "name": name,
                    "attr_text": value.attr_text,
                    "hidden": value.hidden,
                    "local_sheet_id": value.localSheetId,
                }
            )
        calculation = formula_book.calculation
        report = {
            "path": str(source),
            "sheets": sheets,
            "sheet_names": formula_book.sheetnames,
            "active_sheet": formula_book.active.title,
            "defined_names": defined_names,
            "external_links": len(getattr(formula_book, "_external_links", [])),
            "calculation": {
                "mode": calculation.calcMode,
                "full_calc_on_load": calculation.fullCalcOnLoad,
                "force_full_calc": calculation.forceFullCalc,
                "calc_id": calculation.calcId,
            },
            "properties": {
                "title": formula_book.properties.title,
                "subject": formula_book.properties.subject,
                "creator": formula_book.properties.creator,
                "keywords": formula_book.properties.keywords,
                "category": formula_book.properties.category,
                "description": formula_book.properties.description,
            },
            "totals": dict(totals),
            "formula_errors": all_errors,
            "missing_cached_values": all_missing_cached,
            "placeholders": all_placeholders,
        }
        return report
    finally:
        formula_book.close()
        value_book.close()
