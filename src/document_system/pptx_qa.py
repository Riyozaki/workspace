from __future__ import annotations

import math
from collections import Counter
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from pptx import Presentation

from .opc import inspect_ooxml, issue
from .openxml_sdk import validate_with_openxml_sdk
from .pptx_extract import extract_pptx
from .render import render_document
from .util import write_json

IGNORE_OVERLAP_ROLES = {
    "ds-role:panel",
    "ds-role:accent",
    "ds-role:footer",
    "ds-role:source",
    "ds-role:slide-number",
    "ds-role:card",
    "ds-role:metric-card",
    "ds-role:timeline-line",
    "ds-role:timeline-marker",
}


def _inches(value: int) -> float:
    return value / 914400


def _estimated_text_height(shape: Any) -> tuple[float, float]:
    if not getattr(shape, "has_text_frame", False) or not shape.text.strip():
        return 0.0, 0.0
    width_pt = max(10.0, (_inches(shape.width) - 0.12) * 72)
    total_pt = 0.0
    minimum_size = 999.0
    for paragraph in shape.text_frame.paragraphs:
        text = paragraph.text
        sizes = [run.font.size.pt for run in paragraph.runs if run.font.size]
        size = max(sizes, default=18.0)
        minimum_size = min(minimum_size, min(sizes, default=size))
        capacity = max(5, int(width_pt / (size * 0.52)))
        lines = max(1, sum(max(1, math.ceil(len(piece) / capacity)) for piece in text.split("\n")))
        total_pt += lines * size * 1.18 + 4
    return total_pt / 72, minimum_size if minimum_size != 999 else 18.0


def _overlap(first: Any, second: Any) -> float:
    ax1, ay1, ax2, ay2 = first.left, first.top, first.left + first.width, first.top + first.height
    bx1, by1, bx2, by2 = second.left, second.top, second.left + second.width, second.top + second.height
    width = max(0, min(ax2, bx2) - max(ax1, bx1))
    height = max(0, min(ay2, by2) - max(ay1, by1))
    intersection = width * height
    if not intersection:
        return 0.0
    return intersection / max(1, min(first.width * first.height, second.width * second.height))


def validate_pptx(
    path: str | Path,
    *,
    output_report: str | Path | None = None,
    render: bool = False,
    render_dir: str | Path | None = None,
    strict: bool = False,
    expected_slides: int | None = None,
    require_titles: bool = True,
    require_notes: bool = False,
    require_text: Iterable[str] = (),
    allow_placeholders: bool = False,
    original: str | Path | None = None,
    openxml_sdk: bool = True,
    require_openxml_sdk: bool = False,
    timeout: int = 240,
) -> dict[str, Any]:
    source = Path(path)
    package = inspect_ooxml(source)
    deck = extract_pptx(source)
    inherited_issues: list[dict[str, str]] = []
    package_issues = list(package["issues"])
    if original:
        baseline = inspect_ooxml(original)
        signatures = {
            (item.get("severity"), item.get("code"), item.get("message"), item.get("part"))
            for item in baseline["issues"]
        }
        inherited_issues = [
            item
            for item in package_issues
            if (item.get("severity"), item.get("code"), item.get("message"), item.get("part"))
            in signatures
        ]
        package_issues = [item for item in package_issues if item not in inherited_issues]
    issues: list[dict[str, str]] = package_issues

    if expected_slides is not None and deck["totals"]["slides"] != expected_slides:
        issues.append(
            issue(
                "error",
                "slide-count-mismatch",
                f"Expected {expected_slides} slides, found {deck['totals']['slides']}",
            )
        )
    if deck["totals"]["slides"] == 0:
        issues.append(issue("error", "empty-presentation", "Presentation has no slides"))
    if deck["placeholders"] and not allow_placeholders:
        for placeholder in deck["placeholders"][:25]:
            issues.append(
                issue(
                    "error",
                    "placeholder-present",
                    f"Slide {placeholder['slide']} shape {placeholder['shape']} contains {placeholder['value']!r}",
                )
            )
    combined_text = "\n".join(slide["text"] for slide in deck["slides"])
    for text in require_text:
        if text not in combined_text:
            issues.append(issue("error", "required-text-missing", f"Required presentation text not found: {text!r}"))
    for slide in deck["slides"]:
        if require_titles and not slide["title"]:
            issues.append(issue("error", "slide-title-missing", f"Slide {slide['number']} has no identifiable title"))
        if not slide["text"].strip() and not any(shape["type"].endswith("PICTURE (13)") for shape in slide["shapes"]):
            issues.append(issue("warning", "empty-slide", f"Slide {slide['number']} has no text or picture"))
        if require_notes and not slide["notes"]:
            issues.append(issue("error", "speaker-notes-missing", f"Slide {slide['number']} has no speaker notes"))

    presentation = Presentation(source)
    slide_width = presentation.slide_width
    slide_height = presentation.slide_height
    for slide_number, slide in enumerate(presentation.slides, start=1):
        candidates = []
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > slide_width or shape.top + shape.height > slide_height:
                issues.append(issue("error", "shape-out-of-bounds", f"Slide {slide_number}: {shape.name!r} extends outside the slide"))
            estimated_height, minimum_size = _estimated_text_height(shape)
            if estimated_height and estimated_height > _inches(shape.height) * 1.05:
                issues.append(
                    issue(
                        "error",
                        "possible-text-overflow",
                        f"Slide {slide_number}: {shape.name!r} needs about {estimated_height:.2f}in for {_inches(shape.height):.2f}in box",
                    )
                )
            if (
                estimated_height
                and minimum_size < 9
                and shape.name not in {"ds-role:footer", "ds-role:source", "ds-role:slide-number"}
            ):
                issues.append(issue("warning", "small-text", f"Slide {slide_number}: {shape.name!r} uses {minimum_size:.1f}pt text"))
            if shape.name not in IGNORE_OVERLAP_ROLES:
                candidates.append(shape)
        for index, first in enumerate(candidates):
            for second in candidates[index + 1 :]:
                ratio = _overlap(first, second)
                if ratio > 0.3:
                    # Containment of text in a background shape is intentional only for named cards/panels,
                    # which were excluded above. Other large overlap deserves review.
                    issues.append(
                        issue(
                            "warning",
                            "possible-shape-overlap",
                            f"Slide {slide_number}: {first.name!r} and {second.name!r} overlap {ratio:.0%}",
                        )
                    )

    titles = [slide["title"] for slide in deck["slides"]]
    for index in range(1, len(titles)):
        if titles[index] and titles[index] == titles[index - 1]:
            issues.append(issue("warning", "duplicate-adjacent-title", f"Slides {index} and {index + 1} share the same title"))

    schema_report = None
    if openxml_sdk:
        schema_report = validate_with_openxml_sdk(source, timeout=timeout)
        if schema_report.get("available") and schema_report.get("valid") is False:
            issues.append(issue("error", "openxml-sdk-invalid", f"Open XML SDK reported {schema_report.get('error_count', 'one or more')} error(s)"))
        elif require_openxml_sdk and not schema_report.get("available"):
            issues.append(issue("error", "openxml-sdk-unavailable", f"Open XML SDK validation was required: {schema_report.get('reason')}"))

    render_report = None
    if render:
        destination = Path(render_dir or source.parent / f"{source.stem}-render")
        try:
            render_report = render_document(source, destination, timeout=timeout)
            if render_report["page_count"] != deck["totals"]["slides"]:
                issues.append(issue("error", "rendered-slide-count-mismatch", f"Rendered {render_report['page_count']} pages for {deck['totals']['slides']} slides"))
            for page in render_report["summary"]["blank_pages"]:
                issues.append(issue("warning", "likely-blank-slide", f"Rendered slide {page} appears blank"))
            for page in render_report["summary"]["possible_edge_clipping_pages"]:
                rendered_slide = presentation.slides[page - 1]
                intentional = any(
                    shape.name in {"ds-role:panel", "ds-role:decor"}
                    and (
                        shape.left <= 0
                        or shape.top <= 0
                        or shape.left + shape.width >= slide_width
                        or shape.top + shape.height >= slide_height
                    )
                    for shape in rendered_slide.shapes
                )
                try:
                    background_rgb = rendered_slide.background.fill.fore_color.rgb
                    intentional = intentional or (
                        background_rgb is not None
                        and str(background_rgb).upper() not in {"FFFFFF", "F7F9FA"}
                    )
                except (AttributeError, TypeError):
                    pass
                if not intentional:
                    issues.append(
                        issue(
                            "warning",
                            "possible-edge-clipping",
                            f"Rendered slide {page} has ink at the edge",
                        )
                    )
        except Exception as exc:
            issues.append(issue("error", "render-failed", str(exc)))

    counts = Counter(item["severity"] for item in issues)
    passed = counts["error"] == 0 and (not strict or counts["warning"] == 0)
    report = {
        "path": str(source),
        "format": "pptx",
        "passed": passed,
        "strict": strict,
        "summary": {"errors": counts["error"], "warnings": counts["warning"], "info": counts["info"]},
        "issues": issues,
        "package": package,
        "inherited_package_issues": inherited_issues,
        "presentation": deck,
        "schema": schema_report,
        "render": render_report,
    }
    if output_report:
        write_json(output_report, report)
    return report
