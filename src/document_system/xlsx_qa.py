from __future__ import annotations

from collections import Counter
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .opc import inspect_ooxml, issue
from .openxml_sdk import validate_with_openxml_sdk
from .render import render_document
from .util import write_json
from .xlsx_extract import extract_xlsx
from .xlsx_formula import lint_formulas


def validate_xlsx(
    path: str | Path,
    *,
    output_report: str | Path | None = None,
    render: bool = False,
    render_dir: str | Path | None = None,
    strict: bool = False,
    require_sheets: Iterable[str] = (),
    require_cells: Iterable[str] = (),
    require_formulas: bool = False,
    require_recalculated: bool = False,
    libreoffice_compatible: bool = False,
    allow_placeholders: bool = False,
    openxml_sdk: bool = True,
    require_openxml_sdk: bool = False,
    timeout: int = 240,
) -> dict[str, Any]:
    source = Path(path)
    package = inspect_ooxml(source)
    workbook = extract_xlsx(source, include_cells=True)
    issues: list[dict[str, str]] = list(package["issues"])

    if workbook["totals"].get("visible_sheets", 0) == 0:
        issues.append(issue("error", "no-visible-worksheet", "Workbook has no visible worksheet"))
    for sheet in require_sheets:
        if sheet not in workbook["sheet_names"]:
            issues.append(issue("error", "required-sheet-missing", f"Required worksheet is missing: {sheet}"))
    if require_cells:
        check_book = load_workbook(source, data_only=False, read_only=True)
        try:
            for reference in require_cells:
                if "!" not in reference:
                    issues.append(issue("error", "invalid-required-cell", f"Use Sheet!A1 notation: {reference}"))
                    continue
                sheet_name, coordinate = reference.rsplit("!", 1)
                if sheet_name not in check_book.sheetnames or check_book[sheet_name][coordinate].value is None:
                    issues.append(issue("error", "required-cell-empty", f"Required cell is empty or missing: {reference}"))
        finally:
            check_book.close()

    if require_formulas and workbook["totals"].get("formulas", 0) == 0:
        issues.append(issue("error", "formulas-required", "Workbook contains no formulas"))
    for error in workbook["formula_errors"]:
        issues.append(
            issue(
                "error",
                "formula-error",
                f"{error['sheet']}!{error['cell']} evaluates to {error['error']}",
            )
        )
    formula_analysis = lint_formulas(source, libreoffice_compatible=libreoffice_compatible)
    for formula_issue in formula_analysis["issues"]:
        issues.append(
            issue(
                formula_issue["severity"],
                formula_issue["code"],
                f"{formula_issue['cell']}: {formula_issue['message']}",
            )
        )

    missing_cached = workbook["missing_cached_values"]
    if missing_cached:
        severity = "error" if require_recalculated else "warning"
        sample = ", ".join(f"{item['sheet']}!{item['cell']}" for item in missing_cached[:20])
        remainder = len(missing_cached) - 20
        if remainder > 0:
            sample += f" and {remainder} more"
        issues.append(
            issue(
                severity,
                "formula-cache-missing",
                f"Formula cached values are missing: {sample}. Recalculate with LibreOffice/Excel.",
            )
        )
    if workbook["placeholders"] and not allow_placeholders:
        for placeholder in workbook["placeholders"][:25]:
            issues.append(
                issue(
                    "error",
                    "placeholder-present",
                    f"{placeholder['sheet']}!{placeholder['cell']} contains {placeholder['value']!r}",
                )
            )
    calculation = workbook["calculation"]
    if workbook["totals"].get("formulas", 0) and calculation["mode"] == "manual":
        issues.append(issue("error", "manual-calculation", "Workbook with formulas is configured for manual calculation"))
    if workbook["external_links"]:
        issues.append(
            issue(
                "warning",
                "external-workbook-links",
                f"Workbook contains {workbook['external_links']} external link definition(s)",
            )
        )

    schema_report = None
    if openxml_sdk:
        schema_report = validate_with_openxml_sdk(source, timeout=timeout)
        if schema_report.get("available") and schema_report.get("valid") is False:
            issues.append(
                issue(
                    "error",
                    "openxml-sdk-invalid",
                    f"Open XML SDK reported {schema_report.get('error_count', 'one or more')} error(s)",
                )
            )
        elif require_openxml_sdk and not schema_report.get("available"):
            issues.append(
                issue(
                    "error",
                    "openxml-sdk-unavailable",
                    f"Open XML SDK validation was required: {schema_report.get('reason')}",
                )
            )

    render_report = None
    if render:
        destination = Path(render_dir or source.parent / f"{source.stem}-render")
        try:
            render_report = render_document(source, destination, timeout=timeout)
            for page in render_report["summary"]["blank_pages"]:
                issues.append(issue("warning", "likely-blank-page", f"Rendered page {page} appears blank"))
            for page in render_report["summary"]["possible_edge_clipping_pages"]:
                issues.append(issue("warning", "possible-edge-clipping", f"Rendered page {page} has ink at the edge"))
        except Exception as exc:
            issues.append(issue("error", "render-failed", str(exc)))

    counts = Counter(item["severity"] for item in issues)
    passed = counts["error"] == 0 and (not strict or counts["warning"] == 0)
    report = {
        "path": str(source),
        "format": "xlsx",
        "passed": passed,
        "strict": strict,
        "summary": {
            "errors": counts["error"],
            "warnings": counts["warning"],
            "info": counts["info"],
        },
        "issues": issues,
        "package": package,
        "workbook": workbook,
        "formula_analysis": formula_analysis,
        "schema": schema_report,
        "render": render_report,
    }
    if output_report:
        write_json(output_report, report)
    return report
