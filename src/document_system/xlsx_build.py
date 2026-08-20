from __future__ import annotations

import importlib.resources
import json
import re
import zipfile
from copy import copy
from pathlib import Path
from typing import Any

import jsonschema
from lxml import etree
from openpyxl import Workbook
from openpyxl.chart import AreaChart, BarChart, LineChart, PieChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, NamedStyle, PatternFill, Protection, Side
from openpyxl.utils.cell import range_boundaries
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from .util import read_json

DEFAULT_THEME = {
    "primary": "24445C",
    "secondary": "3F6B7D",
    "accent": "C7863B",
    "positive": "26734D",
    "negative": "B54343",
    "text": "24323D",
    "muted": "657681",
    "light": "EEF3F5",
    "font": "Arial",
}

INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")


def _argb(value: str) -> str:
    value = value.strip().lstrip("#").upper()
    if len(value) != 6 or any(char not in "0123456789ABCDEF" for char in value):
        raise ValueError(f"Expected six-digit RGB color, got {value!r}")
    return "FF" + value


def validate_xlsx_spec(spec: dict[str, Any]) -> None:
    schema_path = importlib.resources.files("document_system").joinpath("schemas/xlsx-spec.schema.json")
    schema = json.loads(schema_path.read_text(encoding="utf-8"))
    jsonschema.Draft202012Validator(schema).validate(spec)
    names = [sheet["name"] for sheet in spec["sheets"]]
    folded = [name.casefold() for name in names]
    if len(folded) != len(set(folded)):
        raise ValueError("Worksheet names must be unique ignoring case")
    for name in names:
        if INVALID_SHEET_CHARS.search(name) or name.startswith("'") or name.endswith("'"):
            raise ValueError(f"Invalid worksheet name: {name!r}")
    if not any(sheet.get("state", "visible") == "visible" for sheet in spec["sheets"]):
        raise ValueError("At least one worksheet must be visible")
    active = spec.get("active_sheet")
    if active and active not in names:
        raise ValueError(f"active_sheet does not exist: {active}")
    table_names: list[str] = []
    for sheet in spec["sheets"]:
        table_names.extend(table["name"].casefold() for table in sheet.get("tables", []))
    if len(table_names) != len(set(table_names)):
        raise ValueError("Excel table names must be unique across the workbook")


def _make_named_styles(workbook: Workbook, theme: dict[str, str]) -> None:
    thin = Side(style="thin", color="FFD5DEE3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    font_name = theme["font"]
    styles = [
        NamedStyle(
            name="ds-title",
            font=Font(name=font_name, size=22, bold=True, color=_argb(theme["primary"])),
            alignment=Alignment(horizontal="left", vertical="center"),
        ),
        NamedStyle(
            name="ds-subtitle",
            font=Font(name=font_name, size=12, color=_argb(theme["muted"])),
            alignment=Alignment(horizontal="left", vertical="center"),
        ),
        NamedStyle(
            name="ds-section",
            font=Font(name=font_name, size=13, bold=True, color=_argb(theme["primary"])),
            fill=PatternFill("solid", fgColor=_argb(theme["light"])),
            alignment=Alignment(horizontal="left", vertical="center"),
            border=Border(bottom=Side(style="medium", color=_argb(theme["accent"]))),
        ),
        NamedStyle(
            name="ds-header",
            font=Font(name=font_name, size=10, bold=True, color="FFFFFFFF"),
            fill=PatternFill("solid", fgColor=_argb(theme["primary"])),
            alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
            border=border,
        ),
        NamedStyle(
            name="ds-input",
            font=Font(name=font_name, size=10, color="FF0000FF"),
            fill=PatternFill("solid", fgColor="FFFFF2CC"),
            alignment=Alignment(vertical="center"),
            border=border,
            protection=Protection(locked=False),
        ),
        NamedStyle(
            name="ds-formula",
            font=Font(name=font_name, size=10, color=_argb(theme["text"])),
            alignment=Alignment(vertical="center"),
            border=border,
        ),
        NamedStyle(
            name="ds-text",
            font=Font(name=font_name, size=10, color=_argb(theme["text"])),
            alignment=Alignment(vertical="top", wrap_text=True),
            border=border,
        ),
        NamedStyle(
            name="ds-integer",
            font=Font(name=font_name, size=10, color=_argb(theme["text"])),
            alignment=Alignment(horizontal="right", vertical="center"),
            border=border,
            number_format='#,##0;[Red](#,##0);-',
        ),
        NamedStyle(
            name="ds-decimal",
            font=Font(name=font_name, size=10, color=_argb(theme["text"])),
            alignment=Alignment(horizontal="right", vertical="center"),
            border=border,
            number_format='#,##0.0;[Red](#,##0.0);-',
        ),
        NamedStyle(
            name="ds-currency",
            font=Font(name=font_name, size=10, color=_argb(theme["text"])),
            alignment=Alignment(horizontal="right", vertical="center"),
            border=border,
            number_format='₽#,##0;[Red](₽#,##0);-',
        ),
        NamedStyle(
            name="ds-percent",
            font=Font(name=font_name, size=10, color=_argb(theme["text"])),
            alignment=Alignment(horizontal="right", vertical="center"),
            border=border,
            number_format='0.0%;[Red](0.0%);-',
        ),
        NamedStyle(
            name="ds-date",
            font=Font(name=font_name, size=10, color=_argb(theme["text"])),
            alignment=Alignment(horizontal="center", vertical="center"),
            border=border,
            number_format='dd.mm.yyyy',
        ),
        NamedStyle(
            name="ds-note",
            font=Font(name=font_name, size=9, italic=True, color=_argb(theme["muted"])),
            alignment=Alignment(wrap_text=True, vertical="top"),
        ),
        NamedStyle(
            name="ds-positive",
            font=Font(name=font_name, size=10, bold=True, color=_argb(theme["positive"])),
            fill=PatternFill("solid", fgColor="FFE8F3EC"),
        ),
        NamedStyle(
            name="ds-negative",
            font=Font(name=font_name, size=10, bold=True, color=_argb(theme["negative"])),
            fill=PatternFill("solid", fgColor="FFF8E8E8"),
        ),
    ]
    existing = {style.name for style in workbook._named_styles}
    for style in styles:
        if style.name not in existing:
            workbook.add_named_style(style)


def _apply_cell_style(cell: Any, config: dict[str, Any], theme: dict[str, str]) -> None:
    style_name = config.get("style")
    if style_name:
        cell.style = "ds-" + style_name
    elif config.get("formula"):
        cell.style = "ds-formula"
    if config.get("number_format"):
        cell.number_format = config["number_format"]
    base_font = copy(cell.font)
    cell.font = Font(
        name=base_font.name or theme["font"],
        size=config.get("font_size", base_font.sz),
        bold=config.get("bold", base_font.bold),
        italic=config.get("italic", base_font.italic),
        color=_argb(config["font_color"]) if config.get("font_color") else base_font.color,
    )
    if config.get("fill"):
        cell.fill = PatternFill("solid", fgColor=_argb(config["fill"]))
    alignment = copy(cell.alignment)
    alignment.horizontal = config.get("align", alignment.horizontal)
    alignment.vertical = {"center": "center"}.get(
        config.get("vertical"), config.get("vertical", alignment.vertical)
    )
    alignment.wrap_text = config.get("wrap", alignment.wrap_text)
    cell.alignment = alignment
    if "locked" in config:
        cell.protection = Protection(locked=config["locked"])


def _write_cell(ws: Any, config: dict[str, Any], theme: dict[str, str]) -> None:
    cell = ws[config["cell"].upper()]
    if "formula" in config:
        formula = config["formula"]
        if not formula.startswith("="):
            raise ValueError(f"Formula in {ws.title}!{cell.coordinate} must start with '='")
        cell.value = formula
    elif "value" in config:
        cell.value = config["value"]
    _apply_cell_style(cell, config, theme)
    if config.get("comment"):
        cell.comment = Comment(config["comment"], config.get("comment_author", "Document System"))
    if config.get("hyperlink"):
        cell.hyperlink = config["hyperlink"]
        cell.style = "Hyperlink"


def _write_table(ws: Any, config: dict[str, Any], theme: dict[str, str]) -> None:
    start_col, start_row, _, _ = range_boundaries(config["start_cell"])
    headers = config["headers"]
    for offset, value in enumerate(headers):
        cell = ws.cell(start_row, start_col + offset, value)
        cell.style = "ds-header"
    for row_offset, values in enumerate(config.get("rows", []), start=1):
        if len(values) > len(headers):
            raise ValueError(f"Table {config['name']} row has more values than headers")
        for column_offset, value in enumerate(values):
            cell = ws.cell(start_row + row_offset, start_col + column_offset, value)
            cell.style = "ds-text"
    end_row = start_row + max(1, len(config.get("rows", [])))
    end_col = start_col + len(headers) - 1
    reference = f"{ws.cell(start_row, start_col).coordinate}:{ws.cell(end_row, end_col).coordinate}"
    table = Table(displayName=config["name"], ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name=config.get("style", "TableStyleMedium2"),
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    if config.get("totals_row"):
        table.totalsRowShown = True
    ws.add_table(table)


def _reference(ws: Any, range_string: str) -> Reference:
    min_col, min_row, max_col, max_row = range_boundaries(range_string)
    return Reference(ws, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)


def _write_chart(workbook: Workbook, ws: Any, config: dict[str, Any]) -> None:
    chart_type = config["type"]
    if chart_type in {"bar", "column"}:
        chart = BarChart()
        chart.type = "bar" if chart_type == "bar" else "col"
    elif chart_type == "line":
        chart = LineChart()
    elif chart_type == "pie":
        chart = PieChart()
    elif chart_type == "area":
        chart = AreaChart()
    else:
        raise ValueError(f"Unsupported chart type: {chart_type}")
    data_ws = workbook[config.get("data_sheet", ws.title)]
    chart.add_data(
        _reference(data_ws, config["data_range"]),
        titles_from_data=config.get("titles_from_data", True),
        from_rows=config.get("series_from", "columns") == "rows",
    )
    if config.get("categories_range"):
        chart.set_categories(_reference(data_ws, config["categories_range"]))
    chart.title = config.get("title")
    chart.x_axis.title = config.get("x_title")
    chart.y_axis.title = config.get("y_title")
    chart.style = config.get("style", 10)
    chart.width = config.get("width", 14)
    chart.height = config.get("height", 7.5)
    legend = config.get("legend", "bottom")
    if legend == "none":
        chart.legend = None
    elif chart.legend is not None:
        chart.legend.position = {"top": "t", "bottom": "b", "left": "l", "right": "r"}[legend]
    ws.add_chart(chart, config["anchor"])


def _write_conditional_format(ws: Any, config: dict[str, Any], theme: dict[str, str]) -> None:
    kind = config["type"]
    if kind == "color_scale":
        rule = ColorScaleRule(
            start_type="min",
            start_color=_argb(config.get("min_color", theme["negative"])),
            mid_type="percentile",
            mid_value=50,
            mid_color=_argb(config.get("mid_color", "FFF2CC")),
            end_type="max",
            end_color=_argb(config.get("max_color", theme["positive"])),
        )
    elif kind == "data_bar":
        rule = DataBarRule(
            start_type="min",
            end_type="max",
            color=_argb(config.get("color", theme["secondary"])),
            showValue=True,
        )
    elif kind == "cell":
        fill = PatternFill("solid", fgColor=_argb(config.get("fill", "F8E8E8")))
        font = Font(color=_argb(config.get("font_color", theme["negative"])))
        rule = CellIsRule(
            operator=config.get("operator", "lessThan"),
            formula=config.get("formula", ["0"]),
            fill=fill,
            font=font,
        )
    else:
        raise ValueError(kind)
    ws.conditional_formatting.add(config["range"], rule)


def _write_data_validation(ws: Any, config: dict[str, Any]) -> None:
    validation = DataValidation(
        type=config["type"],
        operator=config.get("operator"),
        formula1=config["formula1"],
        formula2=config.get("formula2"),
        allow_blank=config.get("allow_blank", True),
    )
    validation.promptTitle = config.get("prompt_title")
    validation.prompt = config.get("prompt")
    validation.errorTitle = config.get("error_title")
    validation.error = config.get("error")
    validation.showInputMessage = bool(validation.prompt or validation.promptTitle)
    validation.showErrorMessage = bool(validation.error or validation.errorTitle)
    ws.add_data_validation(validation)
    validation.add(config["range"])


def _configure_sheet(ws: Any, config: dict[str, Any], theme: dict[str, str]) -> None:
    ws.sheet_state = config.get("state", "visible")
    if config.get("tab_color"):
        ws.sheet_properties.tabColor = _argb(config["tab_color"])
    ws.freeze_panes = config.get("freeze_panes")
    ws.sheet_view.showGridLines = config.get("show_gridlines", False)
    ws.sheet_view.zoomScale = config.get("zoom", 90)
    if config.get("auto_filter"):
        ws.auto_filter.ref = config["auto_filter"]
    for merged in config.get("merge_cells", []):
        ws.merge_cells(merged)
    for column in config.get("columns", []):
        dimension = ws.column_dimensions[column["column"].upper()]
        dimension.width = column.get("width", dimension.width)
        dimension.hidden = column.get("hidden", False)
        dimension.outlineLevel = column.get("outline_level", 0)
    for row in config.get("rows", []):
        dimension = ws.row_dimensions[row["row"]]
        dimension.height = row.get("height", dimension.height)
        dimension.hidden = row.get("hidden", False)
        dimension.outlineLevel = row.get("outline_level", 0)
    for table in config.get("tables", []):
        _write_table(ws, table, theme)
    # Explicit cells apply after table materialization so number formats and
    # formulas can intentionally specialize individual table columns/cells.
    for cell in config.get("cells", []):
        _write_cell(ws, cell, theme)
    for validation in config.get("data_validations", []):
        _write_data_validation(ws, validation)
    for conditional in config.get("conditional_formats", []):
        _write_conditional_format(ws, conditional, theme)

    print_config = config.get("print", {})
    if print_config:
        ws.page_setup.orientation = print_config.get("orientation", "portrait")
        ws.page_setup.paperSize = {
            "A4": ws.PAPERSIZE_A4,
            "Letter": ws.PAPERSIZE_LETTER,
        }[print_config.get("paper_size", "A4")]
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = print_config.get("fit_to_width", 1)
        ws.page_setup.fitToHeight = print_config.get("fit_to_height", 0)
        if print_config.get("repeat_rows"):
            ws.print_title_rows = print_config["repeat_rows"]
        if print_config.get("repeat_columns"):
            ws.print_title_cols = print_config["repeat_columns"]
        if print_config.get("print_area"):
            ws.print_area = print_config["print_area"]
        ws.oddHeader.center.text = "&F" if print_config.get("headers_footers", True) else ""
        ws.oddFooter.center.text = "Страница &P из &N" if print_config.get("headers_footers", True) else ""


FONT_CHILD_ORDER = {
    name: index
    for index, name in enumerate(
        (
            "b", "i", "strike", "condense", "extend", "outline", "shadow",
            "u", "vertAlign", "sz", "color", "name", "family", "charset",
            "scheme",
        )
    )
}


def _normalize_openxml_style_order(path: Path) -> None:
    """Normalize openpyxl font child order to the Open XML SDK particle order."""
    with zipfile.ZipFile(path) as source:
        infos = source.infolist()
        values = {info.filename: source.read(info.filename) for info in infos}
    root = etree.fromstring(values["xl/styles.xml"])
    namespace = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    for font in root.xpath("./x:fonts/x:font", namespaces=namespace):
        children = list(font)
        children.sort(
            key=lambda child: FONT_CHILD_ORDER.get(child.tag.rsplit("}", 1)[-1], 999)
        )
        for child in list(font):
            font.remove(child)
        font.extend(children)
    values["xl/styles.xml"] = etree.tostring(
        root, encoding="UTF-8", xml_declaration=True, standalone=True
    )
    temporary = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(temporary, "w") as destination:
        for info in infos:
            destination.writestr(info, values[info.filename])
    temporary.replace(path)


def build_xlsx(spec: dict[str, Any], output_path: str | Path) -> Path:
    validate_xlsx_spec(spec)
    destination = Path(output_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    theme = {**DEFAULT_THEME, **spec.get("theme", {})}
    for key in ("primary", "secondary", "accent", "positive", "negative", "text", "muted", "light"):
        _argb(theme[key])

    workbook = Workbook()
    workbook.remove(workbook.active)
    _make_named_styles(workbook, theme)
    metadata = spec.get("metadata", {})
    properties = workbook.properties
    for source, target in (
        ("title", "title"),
        ("subject", "subject"),
        ("creator", "creator"),
        ("keywords", "keywords"),
        ("category", "category"),
        ("description", "description"),
    ):
        if source in metadata:
            setattr(properties, target, metadata[source])
    workbook.calculation = CalcProperties(
        calcMode="auto",
        fullCalcOnLoad=True,
        forceFullCalc=True,
    )

    for sheet in spec["sheets"]:
        ws = workbook.create_sheet(sheet["name"])
        _configure_sheet(ws, sheet, theme)
    for sheet in spec["sheets"]:
        ws = workbook[sheet["name"]]
        for chart in sheet.get("charts", []):
            _write_chart(workbook, ws, chart)
    if spec.get("active_sheet"):
        workbook.active = workbook.sheetnames.index(spec["active_sheet"])
    workbook.save(destination)
    _normalize_openxml_style_order(destination)
    return destination


def build_xlsx_from_file(spec_path: str | Path, output_path: str | Path) -> Path:
    return build_xlsx(read_json(spec_path), output_path)
