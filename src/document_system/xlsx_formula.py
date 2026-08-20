from __future__ import annotations

import re
from collections import defaultdict, deque
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import range_boundaries

FORBIDDEN_LINEAR_RUNTIME_FUNCTIONS = {
    "XLOOKUP",
    "XMATCH",
    "SORT",
    "SORTBY",
    "FILTER",
    "UNIQUE",
    "SEQUENCE",
}
PREFIX_REQUIRED_FUNCTIONS = {
    "TEXTJOIN",
    "CONCAT",
    "IFS",
    "SWITCH",
    "MAXIFS",
    "MINIFS",
}
VOLATILE_FUNCTIONS = {"INDIRECT", "OFFSET", "RAND", "RANDBETWEEN", "NOW", "TODAY"}
CELL_RANGE_RE = re.compile(
    r"^(?:(?P<sheet>'(?:[^']|'')+'|[^!]+)!)?(?P<start>\$?[A-Za-z]{1,3}\$?[1-9][0-9]*)(?::(?P<end>\$?[A-Za-z]{1,3}\$?[1-9][0-9]*))?$"
)
FUNCTION_RE = re.compile(r"(?<![A-Z0-9_.])(?P<name>_xlfn\.)?(?P<function>[A-Z][A-Z0-9_.]*)\s*\(", re.IGNORECASE)


def _normalize_sheet(value: str) -> str:
    value = value.strip()
    if value.startswith("'") and value.endswith("'"):
        value = value[1:-1].replace("''", "'")
    return value


def _normalize_cell(value: str) -> str:
    return value.replace("$", "").upper()


def _extract_references(formula: str, current_sheet: str, sheet_names: set[str], max_range_cells: int) -> tuple[list[str], list[str]]:
    dependencies: list[str] = []
    unresolved: list[str] = []
    try:
        tokens = Tokenizer(formula).items
    except Exception:
        return [], [formula]
    for token in tokens:
        if token.subtype != "RANGE":
            continue
        raw = token.value.strip()
        if "[" in raw and "]" in raw:
            unresolved.append(raw)
            continue
        match = CELL_RANGE_RE.match(raw)
        if not match:
            unresolved.append(raw)
            continue
        sheet = _normalize_sheet(match.group("sheet")) if match.group("sheet") else current_sheet
        if sheet not in sheet_names:
            unresolved.append(raw)
            continue
        start = _normalize_cell(match.group("start"))
        end = _normalize_cell(match.group("end") or match.group("start"))
        min_col, min_row, max_col, max_row = range_boundaries(f"{start}:{end}")
        count = (max_col - min_col + 1) * (max_row - min_row + 1)
        if count > max_range_cells:
            unresolved.append(f"{sheet}!{start}:{end} ({count} cells)")
            continue
        from openpyxl.utils import get_column_letter

        for row in range(min_row, max_row + 1):
            for column in range(min_col, max_col + 1):
                dependencies.append(f"{sheet}!{get_column_letter(column)}{row}")
    return sorted(set(dependencies)), sorted(set(unresolved))


def build_formula_graph(path: str | Path, *, max_range_cells: int = 10_000) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=False, read_only=False, keep_links=True)
    try:
        sheet_names = set(workbook.sheetnames)
        formulas: dict[str, str] = {}
        dependencies: dict[str, list[str]] = {}
        unresolved: dict[str, list[str]] = {}
        reverse: dict[str, list[str]] = defaultdict(list)
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.data_type != "f" and not (
                        isinstance(cell.value, str) and cell.value.startswith("=")
                    ):
                        continue
                    key = f"{worksheet.title}!{cell.coordinate}"
                    formula = str(cell.value)
                    formulas[key] = formula
                    refs, unknown = _extract_references(
                        formula, worksheet.title, sheet_names, max_range_cells
                    )
                    dependencies[key] = refs
                    if unknown:
                        unresolved[key] = unknown
                    for dependency in refs:
                        reverse[dependency].append(key)

        formula_nodes = set(formulas)
        graph = {
            node: [dependency for dependency in dependencies[node] if dependency in formula_nodes]
            for node in formula_nodes
        }
        cycles: list[list[str]] = []
        visiting: set[str] = set()
        visited: set[str] = set()
        stack: list[str] = []

        def visit(node: str) -> None:
            if node in visiting:
                start = stack.index(node)
                cycle = stack[start:] + [node]
                if cycle not in cycles:
                    cycles.append(cycle)
                return
            if node in visited:
                return
            visiting.add(node)
            stack.append(node)
            for dependency in graph.get(node, []):
                visit(dependency)
            stack.pop()
            visiting.remove(node)
            visited.add(node)

        for node in sorted(formula_nodes):
            visit(node)
        return {
            "path": str(Path(path)),
            "formula_count": len(formulas),
            "formulas": formulas,
            "dependencies": dependencies,
            "dependents": {key: sorted(value) for key, value in reverse.items()},
            "unresolved_references": unresolved,
            "cycles": cycles,
        }
    finally:
        workbook.close()


def lint_formulas(path: str | Path, *, libreoffice_compatible: bool = False) -> dict[str, Any]:
    graph = build_formula_graph(path)
    issues: list[dict[str, str]] = []
    for cell, formula in graph["formulas"].items():
        upper = formula.upper()
        if ";" in formula:
            issues.append(
                {
                    "severity": "warning",
                    "code": "localized-separator",
                    "cell": cell,
                    "message": "Formula uses ';'; OOXML formulas normally use comma separators",
                }
            )
        for match in FUNCTION_RE.finditer(formula):
            function = match.group("function").upper()
            prefixed = bool(match.group("name"))
            if function in FORBIDDEN_LINEAR_RUNTIME_FUNCTIONS:
                issues.append(
                    {
                        "severity": "error" if libreoffice_compatible else "warning",
                        "code": "unsupported-linear-runtime-function",
                        "cell": cell,
                        "message": f"{function} is not safe for the pinned LibreOffice verification route",
                    }
                )
            if function in PREFIX_REQUIRED_FUNCTIONS and not prefixed:
                issues.append(
                    {
                        "severity": "error" if libreoffice_compatible else "warning",
                        "code": "xlfn-prefix-required",
                        "cell": cell,
                        "message": f"{function} should be stored with an _xlfn. prefix for compatibility",
                    }
                )
            if function in VOLATILE_FUNCTIONS:
                issues.append(
                    {
                        "severity": "warning",
                        "code": "volatile-function",
                        "cell": cell,
                        "message": f"{function} is volatile or difficult to audit",
                    }
                )
        if "#REF!" in upper:
            issues.append(
                {
                    "severity": "error",
                    "code": "broken-formula-reference",
                    "cell": cell,
                    "message": "Formula contains #REF!",
                }
            )
    for cell, values in graph["unresolved_references"].items():
        for value in values:
            issues.append(
                {
                    "severity": "warning",
                    "code": "unresolved-or-external-reference",
                    "cell": cell,
                    "message": f"Reference could not be expanded safely: {value}",
                }
            )
    for cycle in graph["cycles"]:
        issues.append(
            {
                "severity": "error",
                "code": "circular-reference",
                "cell": cycle[0],
                "message": "Circular formula chain: " + " -> ".join(cycle),
            }
        )
    return {"graph": graph, "issues": issues}


def trace_formula(path: str | Path, reference: str, *, direction: str = "precedents", max_depth: int = 20) -> dict[str, Any]:
    graph = build_formula_graph(path)
    mapping = graph["dependencies"] if direction == "precedents" else graph["dependents"]
    if direction not in {"precedents", "dependents"}:
        raise ValueError("direction must be precedents or dependents")
    start = reference.replace("$", "")
    queue = deque([(start, 0, None)])
    seen: set[str] = set()
    nodes: list[dict[str, Any]] = []
    edges: list[dict[str, str]] = []
    while queue:
        node, depth, parent = queue.popleft()
        if parent is not None:
            edges.append({"from": parent, "to": node})
        if node in seen or depth > max_depth:
            continue
        seen.add(node)
        nodes.append(
            {
                "cell": node,
                "depth": depth,
                "formula": graph["formulas"].get(node),
                "citation": f"[{node}]",
            }
        )
        for related in mapping.get(node, []):
            queue.append((related, depth + 1, node))
    return {
        "path": str(Path(path)),
        "start": start,
        "direction": direction,
        "nodes": nodes,
        "edges": edges,
        "truncated": any(node["depth"] == max_depth for node in nodes),
    }
