#!/usr/bin/env python3
"""
render.py — turn any deliverable into PDF and PNG page images so it can be
*looked at* before delivery.

WHY THIS MATTERS
A document that validates can still be visually broken: text overflowing its
box, a table running off the page, a chart with unreadable labels, a cover
whose text is the same colour as its background. None of that shows up in XML
validation. The normal way to catch it is `soffice --convert-to pdf`, which
does not exist here. This routes each format through a converter that does.

    docx  -> pandoc -> styled HTML -> Chromium -> PDF -> PNG
    xlsx  -> HTML table -> Chromium -> PDF -> PNG
    pptx  -> slide-accurate HTML -> Chromium -> PDF -> PNG
    html  -> Chromium -> PDF -> PNG
    md    -> pandoc -> HTML -> Chromium -> PDF -> PNG
    pdf   -> PNG directly

The PDF is a *faithful-enough preview for QA*, not a production PDF export —
pandoc reflows a .docx rather than honouring Word's exact pagination, so use it
to catch layout disasters, not to verify a page break lands on line 34.

USAGE
    python tools/render.py report.docx                  # -> .workdir/preview/
    python tools/render.py deck.pptx -o out/ --dpi 120
    python tools/render.py book.xlsx --pages 1-3
    python tools/render.py report.docx --pdf-only
"""

from __future__ import annotations

import argparse
import html as html_mod
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
NODE_DIR = REPO / "tools"
CSS = REPO / "assets" / "preview.css"


# --------------------------------------------------------------------------- helpers
def _fail(msg: str) -> None:
    print(f"error: {msg}", file=sys.stderr)
    raise SystemExit(2)


def _run_node(script: str, *args: str) -> None:
    """Run a node one-liner from tools/ so require() resolves node_modules."""
    proc = subprocess.run(
        ["node", "-e", script, "--", *args],
        cwd=str(NODE_DIR),
        capture_output=True,
        text=True,
    )
    if proc.returncode != 0:
        _fail(f"chromium render failed:\n{proc.stdout}\n{proc.stderr}")


def html_to_pdf(html_path: Path, pdf_path: Path, landscape: bool = False,
                page_size: str = "A4", margin: str = "18mm") -> Path:
    script = """
    const { launch } = require('./js/chromium.js');
    // `node -e` puts the first user argument at argv[1], not argv[2].
    const [htmlPath, pdfPath, landscape, format, margin] = process.argv.slice(1);
    (async () => {
      const browser = await launch();
      try {
        const page = await browser.newPage();
        await page.goto('file://' + htmlPath, { waitUntil: 'networkidle0' });
        await page.evaluate(() => document.fonts.ready);
        const opts = {
          path: pdfPath, printBackground: true,
          landscape: landscape === 'true',
          margin: { top: margin, bottom: margin, left: margin, right: margin },
        };
        if (format.includes('x')) {
          const [w, h] = format.split('x');
          opts.width = w; opts.height = h;
          opts.margin = { top: '0', bottom: '0', left: '0', right: '0' };
        } else { opts.format = format; }
        await page.pdf(opts);
      } finally { await browser.close(); }
    })().catch(e => { console.error(e.message); process.exit(1); });
    """
    _run_node(script, str(html_path.resolve()), str(pdf_path.resolve()),
              "true" if landscape else "false", page_size, margin)
    return pdf_path


def pdf_to_pngs(pdf_path: Path, out_dir: Path, dpi: int = 110,
                pages: list[int] | None = None) -> list[Path]:
    import pypdfium2 as pdfium

    out_dir.mkdir(parents=True, exist_ok=True)
    doc = pdfium.PdfDocument(str(pdf_path))
    try:
        total = len(doc)
        wanted = pages or list(range(1, total + 1))
        width = len(str(total))
        made: list[Path] = []
        for number in wanted:
            if not 1 <= number <= total:
                continue
            page = doc[number - 1]
            image = page.render(scale=dpi / 72).to_pil()
            target = out_dir / f"page-{number:0{width}d}.png"
            image.save(target)
            image.close()
            page.close()
            made.append(target)
        return made
    finally:
        doc.close()


def _wrap_html(body: str, title: str, extra_css: str = "",
               base_css: bool = True) -> str:
    """
    Wrap a body fragment in a full HTML document.

    base_css=False for slide previews: preview.css sets `@page { size: A4 }`,
    and an @page size rule beats the width/height passed to page.pdf(), which
    silently turns a 10x5.62in slide into a portrait A4 page.
    """
    base = CSS.read_text(encoding="utf-8") if (base_css and CSS.is_file()) else ""
    return f"""<!DOCTYPE html>
<html lang="ru"><head><meta charset="utf-8"><title>{html_mod.escape(title)}</title>
<style>{base}
{extra_css}</style></head><body>{body}</body></html>"""


# --------------------------------------------------------------------------- per-format
def from_markdown(src: Path, work: Path) -> tuple[Path, dict]:
    import pypandoc

    body = pypandoc.convert_file(str(src), "html5", extra_args=["--wrap=none"])
    page = work / "input.html"
    page.write_text(_wrap_html(body, src.stem), encoding="utf-8")
    return page, {}


def from_docx(src: Path, work: Path) -> tuple[Path, dict]:
    import pypandoc

    media = work / "media"
    body = pypandoc.convert_file(
        str(src), "html5",
        extra_args=[f"--extract-media={media}", "--wrap=none"],
    )
    page = work / "input.html"
    page.write_text(_wrap_html(body, src.stem), encoding="utf-8")
    return page, {}


def from_xlsx(src: Path, work: Path) -> tuple[Path, dict]:
    """Render each sheet as a table, showing cached values (run recalc first)."""
    import openpyxl
    from openpyxl.utils import get_column_letter

    wb = openpyxl.load_workbook(src, data_only=True)
    chunks: list[str] = []
    missing = 0

    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        rows_html: list[str] = []
        dims = ws.calculate_dimension()
        for row in ws.iter_rows():
            # Row-number gutter, matching the empty <th> in the header — without
            # it every data cell renders one column to the left of its label.
            cells: list[str] = [
                f'<th class="rowhdr">{row[0].row}</th>' if row else "<th></th>"
            ]
            for cell in row:
                value = cell.value
                if value is None:
                    value = ""
                elif isinstance(value, float):
                    value = f"{value:,.2f}".rstrip("0").rstrip(".")
                style = []
                if cell.font and cell.font.bold:
                    style.append("font-weight:700")
                if cell.font and cell.font.color and cell.font.color.rgb and isinstance(cell.font.color.rgb, str):
                    rgb = cell.font.color.rgb[-6:]
                    if rgb != "000000":
                        style.append(f"color:#{rgb}")
                if cell.fill and cell.fill.fgColor and cell.fill.fgColor.rgb and isinstance(cell.fill.fgColor.rgb, str):
                    rgb = cell.fill.fgColor.rgb[-6:]
                    if rgb not in ("000000", "FFFFFF"):
                        style.append(f"background:#{rgb}")
                if isinstance(cell.value, (int, float)):
                    style.append("text-align:right")
                attr = f' style="{";".join(style)}"' if style else ""
                cells.append(f"<td{attr}>{html_mod.escape(str(value))}</td>")
            rows_html.append("<tr>" + "".join(cells) + "</tr>")

        header = "".join(
            f"<th>{get_column_letter(i)}</th>"
            for i in range(1, ws.max_column + 1)
        )
        chunks.append(
            f"<h2>{html_mod.escape(ws.title)} "
            f'<span class="muted">{dims}</span></h2>'
            f'<table class="grid"><thead><tr><th></th>{header}</tr></thead>'
            f"<tbody>{''.join(rows_html)}</tbody></table>"
        )

    # warn if formulas have no cached value
    wb_f = openpyxl.load_workbook(src, data_only=False)
    wb_v = openpyxl.load_workbook(src, data_only=True)
    for sheet in wb_f.sheetnames:
        fs, vs = wb_f[sheet], wb_v[sheet]
        for row in fs.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    if vs[cell.coordinate].value is None:
                        missing += 1

    banner = ""
    if missing:
        banner = (
            f'<p class="warn">{missing} formula cell(s) have no cached value — '
            "run <code>python tools/recalc.py</code> first or this preview "
            "understates the sheet.</p>"
        )

    page = work / "input.html"
    page.write_text(
        _wrap_html(banner + "".join(chunks), src.stem,
                   "table.grid{font-size:9pt}"),
        encoding="utf-8",
    )
    return page, {"landscape": True, "missing_cached": missing}


def _solid_hex(fill) -> str | None:
    """Return RRGGBB for a solid fill, or None for anything else."""
    try:
        from pptx.enum.dml import MSO_FILL

        if fill.type != MSO_FILL.SOLID:
            return None
        colour = fill.fore_color
        if colour.type is not None and colour.rgb is not None:
            return str(colour.rgb)
    except Exception:
        pass
    return None


def _slide_background(slide) -> str | None:
    """Slide background, falling back to its layout's."""
    for source in (slide, slide.slide_layout):
        try:
            hex_value = _solid_hex(source.background.fill)
        except Exception:
            hex_value = None
        if hex_value:
            return hex_value
    return None


def from_pptx(src: Path, work: Path) -> tuple[Path, dict]:
    """
    Lay each slide out at true proportions with positioned, coloured boxes.

    This is a QA proxy, not a renderer: it reproduces geometry, fills, and text
    colour so you can catch invisible text, overflow, and collisions. Charts and
    pictures are drawn as labelled placeholders at the right position and size —
    enough to see that a chart is present and correctly placed, not what it
    plots. Verify chart *content* from the generating script or the data.
    """
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation(str(src))
    sw, sh = prs.slide_width, prs.slide_height
    slide_w_in, slide_h_in = Emu(sw).inches, Emu(sh).inches
    slides_html: list[str] = []

    for index, slide in enumerate(prs.slides, 1):
        bg = _slide_background(slide)
        boxes: list[str] = []

        for shape in slide.shapes:
            if shape.left is None or shape.top is None:
                continue
            left = Emu(shape.left).inches / slide_w_in * 100
            top = Emu(shape.top).inches / slide_h_in * 100
            width = Emu(shape.width or 0).inches / slide_w_in * 100
            height = Emu(shape.height or 0).inches / slide_h_in * 100
            geom = (
                f"left:{left:.2f}%;top:{top:.2f}%;"
                f"width:{width:.2f}%;height:{height:.2f}%"
            )

            if getattr(shape, "has_chart", False):
                boxes.append(
                    f'<div class="ph chart" style="{geom}">'
                    f"chart: {html_mod.escape(shape.chart.chart_type.name.lower())}</div>"
                )
                continue
            if shape.shape_type == 13:  # PICTURE
                boxes.append(f'<div class="ph pic" style="{geom}">image</div>')
                continue
            if not getattr(shape, "has_text_frame", False):
                continue

            shape_fill = None
            if hasattr(shape, "fill"):
                shape_fill = _solid_hex(shape.fill)
            fill_css = f";background:#{shape_fill}" if shape_fill else ""

            paras: list[str] = []
            for para in shape.text_frame.paragraphs:
                align = {1: "center", 2: "right", 3: "justify"}.get(
                    para.alignment.value if para.alignment is not None else None, "left"
                )
                runs: list[str] = []
                for run in para.runs:
                    css: list[str] = []
                    font = run.font
                    if font.bold:
                        css.append("font-weight:700")
                    if font.italic:
                        css.append("font-style:italic")
                    if font.size:
                        css.append(f"font-size:{font.size.pt:.1f}pt")
                    try:
                        if font.color is not None and font.color.type is not None and font.color.rgb:
                            css.append(f"color:#{font.color.rgb}")
                    except Exception:
                        pass
                    style = f' style="{";".join(css)}"' if css else ""
                    runs.append(f"<span{style}>{html_mod.escape(run.text)}</span>")
                if runs:
                    paras.append(f'<p style="text-align:{align}">{"".join(runs)}</p>')

            if paras:
                boxes.append(f'<div class="tb" style="{geom}{fill_css}">{"".join(paras)}</div>')

        style = f' style="background:#{bg}"' if bg else ""
        note = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            note = (
                '<div class="notes"><b>Notes:</b> '
                + html_mod.escape(slide.notes_slide.notes_text_frame.text.strip())
                + "</div>"
            )
        slides_html.append(
            f'<div class="slide"{style}><div class="num">{index}</div>'
            f'{"".join(boxes)}</div>{note}'
        )

    css = """
    @page { margin: 0; }
    body{margin:0;background:#fff}
    .slide{position:relative;width:100%;height:100vh;overflow:hidden;
           page-break-after:always;background:#fff}
    .tb{position:absolute;overflow:hidden;font-size:12pt;line-height:1.25;
        display:flex;flex-direction:column;justify-content:center}
    .tb p{margin:0 0 .2em}
    .ph{position:absolute;border:1px dashed #98a2b3;color:#667085;font-size:10pt;
        display:flex;align-items:center;justify-content:center;background:#f7f9fb}
    .num{position:absolute;right:6px;bottom:4px;font-size:8pt;color:#98a2b3;z-index:9}
    .notes{font-size:9pt;color:#475467;padding:4mm;page-break-after:always}
    """

    page = work / "input.html"
    page.write_text(
        _wrap_html("".join(slides_html), src.stem, css, base_css=False),
        encoding="utf-8",
    )
    return page, {"page_size": f"{slide_w_in:.2f}inx{slide_h_in:.2f}in", "no_base_css": True}


HANDLERS = {
    ".docx": from_docx, ".dotx": from_docx,
    ".xlsx": from_xlsx, ".xlsm": from_xlsx,
    ".pptx": from_pptx, ".potx": from_pptx,
    ".md": from_markdown, ".markdown": from_markdown,
}


def parse_pages(spec: str | None) -> list[int] | None:
    if not spec:
        return None
    out: list[int] = []
    for part in spec.split(","):
        part = part.strip()
        if "-" in part:
            a, b = part.split("-", 1)
            out.extend(range(int(a), int(b) + 1))
        elif part:
            out.append(int(part))
    return out or None


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("source", type=Path)
    ap.add_argument("-o", "--out", type=Path, default=REPO / ".workdir" / "preview")
    ap.add_argument("--dpi", type=int, default=110)
    ap.add_argument("--pages", help="e.g. 1-3,7")
    ap.add_argument("--pdf-only", action="store_true")
    ap.add_argument("--json", action="store_true")
    args = ap.parse_args()

    src: Path = args.source
    if not src.is_file():
        _fail(f"{src} not found")

    out: Path = args.out
    out.mkdir(parents=True, exist_ok=True)
    suffix = src.suffix.lower()
    pdf_path = out / f"{src.stem}.pdf"
    meta: dict = {}

    with tempfile.TemporaryDirectory() as tmp:
        work = Path(tmp)
        if suffix == ".pdf":
            shutil.copy2(src, pdf_path)
        elif suffix in (".html", ".htm"):
            html_to_pdf(src, pdf_path)
        elif suffix in HANDLERS:
            page, meta = HANDLERS[suffix](src, work)
            html_to_pdf(
                page, pdf_path,
                landscape=bool(meta.get("landscape")),
                page_size=meta.get("page_size", "A4"),
            )
        else:
            _fail(f"unsupported format: {suffix}")

    images: list[Path] = []
    if not args.pdf_only:
        images = pdf_to_pngs(pdf_path, out, dpi=args.dpi, pages=parse_pages(args.pages))

    report = {
        "source": str(src),
        "pdf": str(pdf_path),
        "images": [str(p) for p in images],
        "page_count": len(images),
        **{k: v for k, v in meta.items() if k == "missing_cached"},
    }
    if args.json:
        print(json.dumps(report, indent=2, ensure_ascii=False))
    else:
        print(f"pdf   {pdf_path}")
        for image in images:
            print(f"png   {image}")
        if meta.get("missing_cached"):
            print(f"  ! {meta['missing_cached']} formula cells lack cached values — run recalc.py")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
