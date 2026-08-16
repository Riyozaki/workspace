#!/usr/bin/env python3
"""
validate.py — check that a generated file is actually a valid, sane document.

Two layers:

1. STRUCTURAL — is the OOXML well-formed and complete? A .docx/.xlsx/.pptx is a
   ZIP of XML parts wired together by relationship files. Corruption here is the
   difference between a file that opens and one where Word says "unreadable
   content". We check the zip, the content types, every relationship target, and
   parse each XML part.

2. EDITORIAL — is it fit to send? Leftover `TODO`, `[Company Name]`, `Lorem
   ipsum`, empty headings, formula cells with no cached value. These pass every
   schema check and still embarrass you.

USAGE
    python tools/validate.py report.docx
    python tools/validate.py deck.pptx --json
    python tools/validate.py book.xlsx --strict     # warnings become failures
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from pathlib import Path

OOXML = {".docx", ".dotx", ".xlsx", ".xlsm", ".xltx", ".pptx", ".potx"}

PLACEHOLDER_PATTERNS = [
    (r"\bTODO\b", "TODO marker"),
    (r"\bFIXME\b", "FIXME marker"),
    (r"\bTBD\b", "TBD marker"),
    (r"\bXXX+\b", "XXX marker"),
    (r"\bLorem ipsum\b", "lorem ipsum filler"),
    (r"\[Company Name\]", "unfilled [Company Name]"),
    (r"\[Insert[^\]]*\]", "unfilled [Insert ...]"),
    (r"\[Your [^\]]*\]", "unfilled [Your ...]"),
    (r"\{\{[^}]+\}\}", "unrendered {{template}} tag"),
    (r"Наименование организации", "unfilled Russian placeholder"),
    (r"\bЗАГЛУШКА\b", "Russian placeholder marker"),
]

REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"


class Report:
    def __init__(self, path: Path):
        self.path = path
        self.errors: list[str] = []
        self.warnings: list[str] = []
        self.info: dict[str, object] = {}

    def error(self, msg: str) -> None:
        self.errors.append(msg)

    def warn(self, msg: str) -> None:
        self.warnings.append(msg)

    @property
    def ok(self) -> bool:
        return not self.errors

    def to_dict(self) -> dict:
        return {
            "file": str(self.path),
            "status": "pass" if self.ok else "fail",
            "errors": self.errors,
            "warnings": self.warnings,
            **self.info,
        }


# ------------------------------------------------------------------ structural
def check_zip(path: Path, rep: Report) -> zipfile.ZipFile | None:
    try:
        zf = zipfile.ZipFile(path)
    except zipfile.BadZipFile as exc:
        rep.error(f"not a valid ZIP container: {exc}")
        return None
    bad = zf.testzip()
    if bad:
        rep.error(f"corrupt entry in archive: {bad}")
    if "[Content_Types].xml" not in zf.namelist():
        rep.error("missing [Content_Types].xml — the package is not valid OOXML")
    return zf


def check_xml_parts(zf: zipfile.ZipFile, rep: Report) -> None:
    from lxml import etree

    parsed = 0
    for name in zf.namelist():
        if not name.endswith((".xml", ".rels")):
            continue
        try:
            etree.fromstring(zf.read(name))
            parsed += 1
        except etree.XMLSyntaxError as exc:
            rep.error(f"{name}: malformed XML — {exc}")
    rep.info["xml_parts"] = parsed


def check_relationships(zf: zipfile.ZipFile, rep: Report) -> None:
    """Every internal relationship target must exist in the package."""
    from lxml import etree

    names = set(zf.namelist())
    dangling = 0
    for rels_name in [n for n in names if n.endswith(".rels")]:
        try:
            root = etree.fromstring(zf.read(rels_name))
        except etree.XMLSyntaxError:
            continue  # already reported
        base = rels_name.rsplit("_rels/", 1)[0]
        for rel in root.findall(f"{REL_NS}Relationship"):
            if rel.get("TargetMode") == "External":
                continue
            target = (rel.get("Target") or "").split("#")[0]
            if not target or target.startswith(("http://", "https://", "mailto:")):
                continue
            if target.startswith("/"):
                resolved = target.lstrip("/")
            else:
                parts = (base + target).split("/")
                stack: list[str] = []
                for part in parts:
                    if part == "..":
                        if stack:
                            stack.pop()
                    elif part not in ("", "."):
                        stack.append(part)
                resolved = "/".join(stack)
            if resolved not in names:
                dangling += 1
                rep.error(f"{rels_name}: relationship points at missing part '{target}'")
    rep.info["dangling_relationships"] = dangling


# ------------------------------------------------------------------- editorial
def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix in (".docx", ".dotx"):
            import docx

            document = docx.Document(str(path))
            chunks = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    chunks.extend(cell.text for cell in row.cells)
            for section in document.sections:
                for part in (section.header, section.footer):
                    chunks.extend(p.text for p in part.paragraphs)
            return "\n".join(chunks)

        if suffix in (".pptx", ".potx"):
            from pptx import Presentation

            prs = Presentation(str(path))
            chunks = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False):
                        chunks.append(shape.text_frame.text)
            return "\n".join(chunks)

        if suffix in (".xlsx", ".xlsm", ".xltx"):
            import openpyxl

            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
            chunks = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    chunks.extend(str(v) for v in row if isinstance(v, str))
            wb.close()
            return "\n".join(chunks)
    except Exception:
        return ""
    return ""


def check_placeholders(text: str, rep: Report) -> None:
    for pattern, label in PLACEHOLDER_PATTERNS:
        hits = re.findall(pattern, text, flags=re.IGNORECASE)
        if hits:
            sample = str(hits[0])[:60]
            rep.warn(f"{label} present ({len(hits)}x), e.g. “{sample}”")


def check_docx(path: Path, rep: Report) -> None:
    import docx

    document = docx.Document(str(path))
    paragraphs = document.paragraphs
    rep.info["paragraphs"] = len(paragraphs)
    rep.info["tables"] = len(document.tables)

    empty_headings = [
        p.text for p in paragraphs
        if p.style is not None and p.style.name.startswith("Heading") and not p.text.strip()
    ]
    if empty_headings:
        rep.warn(f"{len(empty_headings)} empty heading paragraph(s)")

    words = sum(len(p.text.split()) for p in paragraphs)
    rep.info["word_count"] = words
    if words == 0:
        rep.error("document contains no body text")

    has_page_number = False
    for section in document.sections:
        for part in (section.footer, section.header):
            for paragraph in part.paragraphs:
                if "PAGE" in paragraph._p.xml:
                    has_page_number = True
    if words > 400 and not has_page_number:
        rep.warn("no page-number field in any header/footer")


# Valid Excel functions the `formulas` engine cannot evaluate. recalc.py leaves
# these cells uncached on purpose, so their absence is not a defect.
ENGINE_GAPS = {
    "SUBTOTAL", "AGGREGATE", "XLOOKUP", "XMATCH", "FILTER", "SORTBY",
    "TEXTSPLIT", "LAMBDA", "LET", "TEXTJOIN", "IFS", "SWITCH",
}


def check_xlsx(path: Path, rep: Report) -> None:
    import openpyxl

    formulas_no_cache = 0
    engine_gap_cells = 0
    error_cells: list[str] = []
    wb_f = openpyxl.load_workbook(path, data_only=False)
    wb_v = openpyxl.load_workbook(path, data_only=True)
    for sheet in wb_f.sheetnames:
        fs, vs = wb_f[sheet], wb_v[sheet]
        for row in fs.iter_rows():
            for cell in row:
                cached = vs[cell.coordinate].value
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    if cached is None:
                        expr = cell.value.upper()
                        if any(fn + "(" in expr for fn in ENGINE_GAPS):
                            # recalc.py deliberately leaves these to Excel.
                            engine_gap_cells += 1
                        else:
                            formulas_no_cache += 1
                if isinstance(cached, str) and cached.startswith("#") and cached.endswith(("!", "?", "A")):
                    error_cells.append(f"{sheet}!{cell.coordinate}={cached}")
    rep.info["sheets"] = wb_f.sheetnames
    rep.info["formulas_without_cached_value"] = formulas_no_cache
    if engine_gap_cells:
        rep.info["cells_left_to_excel"] = engine_gap_cells
    if formulas_no_cache:
        rep.warn(
            f"{formulas_no_cache} formula cell(s) have no cached value — "
            "run tools/recalc.py or they read as empty everywhere but Excel"
        )
    for cell in error_cells[:20]:
        rep.error(f"formula error {cell}")
    wb_f.close()
    wb_v.close()


def check_pptx(path: Path, rep: Report) -> None:
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation(str(path))
    slide_w, slide_h = Emu(prs.slide_width).inches, Emu(prs.slide_height).inches
    rep.info["slides"] = len(prs.slides)
    rep.info["slide_size_in"] = f"{slide_w:.2f}x{slide_h:.2f}"

    if not len(prs.slides):
        rep.error("presentation has no slides")

    for index, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.left is None or shape.top is None:
                continue
            left, top = Emu(shape.left).inches, Emu(shape.top).inches
            width = Emu(shape.width or 0).inches
            height = Emu(shape.height or 0).inches
            if left < -0.05 or top < -0.05:
                rep.warn(f"slide {index}: '{shape.name}' starts off-canvas ({left:.2f}, {top:.2f}in)")
            if left + width > slide_w + 0.05 or top + height > slide_h + 0.05:
                rep.warn(
                    f"slide {index}: '{shape.name}' overflows the slide "
                    f"(ends at {left + width:.2f}, {top + height:.2f}in)"
                )


def check_pdf(path: Path, rep: Report) -> None:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    rep.info["pages"] = len(reader.pages)
    if not len(reader.pages):
        rep.error("PDF has no pages")
    text = "".join((page.extract_text() or "") for page in reader.pages[:20])
    rep.info["extractable_text"] = len(text)
    if len(reader.pages) and not text.strip():
        rep.warn("no extractable text — the PDF may be image-only (bad for search/a11y)")
    check_placeholders(text, rep)


CHECKERS = {
    ".docx": check_docx, ".dotx": check_docx,
    ".xlsx": check_xlsx, ".xlsm": check_xlsx, ".xltx": check_xlsx,
    ".pptx": check_pptx, ".potx": check_pptx,
}


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("files", nargs="+", type=Path)
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--strict", action="store_true", help="treat warnings as failures")
    args = ap.parse_args()

    reports: list[Report] = []
    for path in args.files:
        rep = Report(path)
        if not path.is_file():
            rep.error("file not found")
            reports.append(rep)
            continue

        suffix = path.suffix.lower()
        rep.info["size_bytes"] = path.stat().st_size
        if path.stat().st_size == 0:
            rep.error("file is empty")

        if suffix in OOXML:
            zf = check_zip(path, rep)
            if zf:
                check_xml_parts(zf, rep)
                check_relationships(zf, rep)
                zf.close()
            if rep.ok and suffix in CHECKERS:
                try:
                    CHECKERS[suffix](path, rep)
                except Exception as exc:  # noqa: BLE001
                    rep.error(f"could not inspect content: {exc}")
            check_placeholders(extract_text(path), rep)
        elif suffix == ".pdf":
            try:
                check_pdf(path, rep)
            except Exception as exc:  # noqa: BLE001
                rep.error(f"could not read PDF: {exc}")
        else:
            rep.warn(f"no validator for '{suffix}' — checked existence only")

        reports.append(rep)

    failed = any(not r.ok or (args.strict and r.warnings) for r in reports)

    if args.json:
        print(json.dumps([r.to_dict() for r in reports], indent=2, ensure_ascii=False))
    else:
        for rep in reports:
            mark = "PASS" if rep.ok else "FAIL"
            colour = "\033[32m" if rep.ok else "\033[31m"
            print(f"{colour}{mark}\033[0m  {rep.path.name}")
            for key, value in rep.info.items():
                print(f"        {key}: {value}")
            for err in rep.errors:
                print(f"  \033[31m✗\033[0m {err}")
            for warn in rep.warnings:
                print(f"  \033[33m!\033[0m {warn}")
            print()

    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
