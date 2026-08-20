from __future__ import annotations

from openpyxl import load_workbook

from document_system.tabular import import_tabular_to_xlsx
from document_system.xlsx_qa import validate_xlsx


def test_imports_and_normalizes_ragged_csv(tmp_path):
    source = tmp_path / "messy.csv"
    source.write_text(
        "Name,Value,value,\n"
        "Alpha,10,001,\n"
        "\n"
        "Beta,20.5,002\n"
        "Gamma,true,003,extra\n",
        encoding="utf-8",
    )
    output = tmp_path / "clean.xlsx"
    report = import_tabular_to_xlsx(source, output)
    assert report["delimiter"] == ","
    assert report["blank_rows_removed"] == 1
    assert report["headers"] == ["Name", "Value", "value_2", "Column_4"]
    assert report["ragged_rows"] == [{"row": 4, "columns": 3, "expected": 4}]
    workbook = load_workbook(output)
    try:
        sheet = workbook["Data"]
        assert sheet["B2"].value == 10
        assert sheet["C2"].value == "001"
        assert sheet["B4"].value is True
        assert "ImportedData" in sheet.tables
        assert sheet.freeze_panes == "A2"
    finally:
        workbook.close()
    assert validate_xlsx(output)["passed"] is True
