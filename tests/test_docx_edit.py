from __future__ import annotations

import hashlib
import zipfile

from docx import Document

from document_system.diffing import diff_docx
from document_system.docx_edit import apply_edit_plan
from document_system.docx_extract import extract_docx
from document_system.docx_revisions import apply_revision_view, verify_tracked_text
from document_system.opc import inspect_ooxml


def _split_run_document(path):
    document = Document()
    document.core_properties.title = "Template"
    paragraph = document.add_paragraph()
    paragraph.add_run("Client: {{CLI").bold = True
    paragraph.add_run("ENT_NAME}}")
    paragraph.add_run(". Payment is due within ")
    paragraph.add_run("30 calendar").italic = True
    paragraph.add_run(" days after acceptance.")
    document.add_paragraph("This paragraph must remain unchanged.")
    document.save(path)


def _hash_parts(path):
    with zipfile.ZipFile(path) as archive:
        return {
            name: hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if not name.endswith("/")
        }


def test_replaces_placeholder_split_across_runs(tmp_path):
    source = tmp_path / "template.docx"
    output = tmp_path / "filled.docx"
    _split_run_document(source)
    report = apply_edit_plan(
        source,
        output,
        {
            "replacements": [
                {
                    "find": "{{CLIENT_NAME}}",
                    "replace": "Northern Systems LLC",
                    "expected": 1,
                }
            ]
        },
    )
    assert report["replacements"][0]["count"] == 1
    content = extract_docx(output)["text"]
    assert "Client: Northern Systems LLC." in content
    assert "{{CLIENT_NAME}}" not in content
    package = inspect_ooxml(output)
    assert package["summary"]["errors"] == 0, package["issues"]


def test_tracked_change_and_comment_preserve_unrelated_parts(tmp_path):
    source = tmp_path / "contract.docx"
    output = tmp_path / "reviewed.docx"
    _split_run_document(source)
    before = _hash_parts(source)
    report = apply_edit_plan(
        source,
        output,
        {
            "author": "Legal Review",
            "track_changes": True,
            "comments": [
                {
                    "target": "30 calendar days",
                    "text": "Confirm the shorter business-day period.",
                    "phase": "before",
                    "initials": "LR",
                }
            ],
            "replacements": [
                {
                    "find": "30 calendar days",
                    "replace": "20 business days",
                    "expected": 1,
                }
            ],
        },
    )
    assert report["comments"][0]["id"] == 0
    accepted = extract_docx(output, revisions="accept")
    rejected = extract_docx(output, revisions="reject")
    assert "20 business days" in accepted["text"]
    assert "30 calendar days" not in accepted["text"]
    assert "30 calendar days" in rejected["text"]
    assert "20 business days" not in rejected["text"]
    assert accepted["inventory"]["insertions"] == 1
    assert accepted["inventory"]["deletions"] == 1
    assert accepted["inventory"]["comments"] == 1
    assert accepted["settings"]["track_revisions"] is True

    package = inspect_ooxml(output)
    assert package["summary"]["errors"] == 0, package["issues"]
    with zipfile.ZipFile(output) as archive:
        document_xml = archive.read("word/document.xml")
        assert b"commentRangeStart" in document_xml
        assert b"commentRangeEnd" in document_xml
        assert b"commentReference" in document_xml
        assert b"<w:ins" in document_xml
        assert b"<w:del" in document_xml

    after = _hash_parts(output)
    allowed_changed = {
        "word/document.xml",
        "word/settings.xml",
        "word/comments.xml",
        "word/_rels/document.xml.rels",
        "[Content_Types].xml",
    }
    for name in before.keys() & after.keys():
        if name not in allowed_changed:
            assert before[name] == after[name], name

    difference = diff_docx(source, output)
    assert difference["text"]["changed"] is True
    assert set(difference["package"]["added_parts"]) == {"word/comments.xml"}


def test_accept_reject_and_tracked_completeness(tmp_path):
    source = tmp_path / "source.docx"
    reviewed = tmp_path / "reviewed.docx"
    accepted = tmp_path / "accepted.docx"
    rejected = tmp_path / "rejected.docx"
    _split_run_document(source)
    apply_edit_plan(
        source,
        reviewed,
        {
            "track_changes": True,
            "replacements": [
                {
                    "find": "30 calendar days",
                    "replace": "20 business days",
                    "expected": 1,
                }
            ],
        },
    )
    proof = verify_tracked_text(source, reviewed)
    assert proof["passed"] is True
    apply_revision_view(reviewed, accepted, mode="accept")
    apply_revision_view(reviewed, rejected, mode="reject")
    assert "20 business days" in extract_docx(accepted)["text"]
    assert "30 calendar days" in extract_docx(rejected)["text"]
    assert extract_docx(accepted)["inventory"]["insertions"] == 0
    assert extract_docx(rejected)["inventory"]["deletions"] == 0

    untracked = tmp_path / "untracked.docx"
    apply_edit_plan(
        source,
        untracked,
        {
            "replacements": [
                {"find": "30 calendar days", "replace": "10 days", "expected": 1}
            ]
        },
    )
    assert verify_tracked_text(source, untracked)["passed"] is False


def test_expected_count_mismatch_is_atomic(tmp_path):
    source = tmp_path / "template.docx"
    output = tmp_path / "should-not-exist.docx"
    _split_run_document(source)
    import pytest

    from document_system.docx_edit import EditError

    with pytest.raises(EditError, match="expected 2"):
        apply_edit_plan(
            source,
            output,
            {
                "replacements": [
                    {"find": "{{CLIENT_NAME}}", "replace": "X", "expected": 2}
                ]
            },
        )
    assert not output.exists()
