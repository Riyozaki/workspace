from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from document_system.opc import inspect_ooxml
from document_system.xlsx_build import build_xlsx
from document_system.xlsx_edit import XlsxEditError, apply_xlsx_edit_plan
from document_system.xlsx_extract import extract_xlsx
from document_system.xlsx_qa import validate_xlsx


@pytest.fixture
def dashboard_spec():
    path = Path(__file__).resolve().parents[1] / "examples/xlsx/operational-dashboard.json"
    return json.loads(path.read_text(encoding="utf-8"))


def _hash_parts(path):
    with zipfile.ZipFile(path) as archive:
        return {
            name: hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if not name.endswith("/")
        }


def test_builds_structured_workbook(tmp_path, dashboard_spec):
    output = build_xlsx(dashboard_spec, tmp_path / "dashboard.xlsx")
    package = inspect_ooxml(output)
    assert package["summary"]["errors"] == 0, package["issues"]
    extracted = extract_xlsx(output)
    assert extracted["active_sheet"] == "Панель"
    assert extracted["totals"]["formulas"] == 4
    assert extracted["totals"]["charts"] == 1
    assert extracted["totals"]["tables"] == 1
    assert extracted["totals"]["visible_sheets"] == 3
    assert extracted["calculation"]["mode"] == "auto"
    assert extracted["calculation"]["full_calc_on_load"] is True
    assert len(extracted["missing_cached_values"]) == 4

    workbook = load_workbook(output, data_only=False)
    try:
        assert workbook["Данные"]["B2"].number_format.startswith("₽")
        assert workbook["Данные"]["C2"].number_format.startswith("0.0%")
        assert workbook["Допущения"]["B2"].protection.locked is False
        assert len(workbook["Панель"]._charts) == 1
        assert "QuarterlyData" in workbook["Данные"].tables
    finally:
        workbook.close()

    report = validate_xlsx(
        output,
        require_sheets=["Панель", "Данные", "Допущения"],
        require_cells=["Панель!B5", "Допущения!B2"],
        require_formulas=True,
    )
    assert report["passed"] is True
    assert any(item["code"] == "formula-cache-missing" for item in report["issues"])


def test_require_recalculation_is_a_hard_gate(tmp_path, dashboard_spec):
    output = build_xlsx(dashboard_spec, tmp_path / "dashboard.xlsx")
    report = validate_xlsx(output, require_formulas=True, require_recalculated=True)
    assert report["passed"] is False
    assert any(
        item["code"] == "formula-cache-missing" and item["severity"] == "error"
        for item in report["issues"]
    )


def test_direct_edit_preserves_unrelated_parts(tmp_path, dashboard_spec):
    source = build_xlsx(dashboard_spec, tmp_path / "dashboard.xlsx")
    output = tmp_path / "scenario.xlsx"
    before = _hash_parts(source)
    result = apply_xlsx_edit_plan(
        source,
        output,
        {
            "updates": [
                {"sheet": "Допущения", "cell": "B2", "expected": 0.22, "value": 0.24},
                {"sheet": "Допущения", "cell": "B3", "expected": "Базовый", "value": "Консервативный"},
                {
                    "sheet": "Панель",
                    "cell": "B5",
                    "expected": "=SUM(Данные!B2:B5)",
                    "formula": "=SUM(Данные!B2:B5)*0.95",
                },
            ],
            "metadata": {"description": "Scenario changed"},
        },
    )
    assert result["formula_changed"] is True
    extracted = extract_xlsx(output)
    panel = next(sheet for sheet in extracted["sheets"] if sheet["name"] == "Панель")
    assumptions = next(sheet for sheet in extracted["sheets"] if sheet["name"] == "Допущения")
    assert next(item for item in panel["formulas"] if item["cell"] == "B5")["formula"].endswith("*0.95")
    values = {item["cell"]: item["value"] for item in assumptions["cells"]}
    assert values["B2"] == 0.24
    assert values["B3"] == "Консервативный"

    after = _hash_parts(output)
    allowed = {
        "docProps/core.xml",
        "xl/workbook.xml",
        "xl/worksheets/sheet1.xml",
        "xl/worksheets/sheet3.xml",
    }
    for name in before.keys() & after.keys():
        if name not in allowed:
            assert before[name] == after[name], name
    package = inspect_ooxml(output)
    assert package["summary"]["errors"] == 0, package["issues"]


def test_broken_formula_fails_validation(tmp_path, dashboard_spec):
    source = build_xlsx(dashboard_spec, tmp_path / "dashboard.xlsx")
    output = tmp_path / "broken.xlsx"
    apply_xlsx_edit_plan(
        source,
        output,
        {
            "updates": [
                {"sheet": "Панель", "cell": "B5", "formula": "=#REF!"}
            ]
        },
    )
    report = validate_xlsx(output)
    assert report["passed"] is False
    assert any(item["code"] == "broken-formula-reference" for item in report["issues"])


def test_expected_value_mismatch_is_atomic(tmp_path, dashboard_spec):
    source = build_xlsx(dashboard_spec, tmp_path / "dashboard.xlsx")
    output = tmp_path / "not-created.xlsx"
    with pytest.raises(XlsxEditError, match="expected"):
        apply_xlsx_edit_plan(
            source,
            output,
            {"updates": [{"sheet": "Допущения", "cell": "B2", "expected": 0.99, "value": 0.24}]},
        )
    assert not output.exists()
