from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

from document_system.opc import inspect_ooxml
from document_system.pptx_build import build_pptx
from document_system.pptx_edit import PptxEditError, apply_pptx_edit_plan
from document_system.pptx_extract import extract_pptx
from document_system.pptx_qa import validate_pptx
from document_system.pptx_structure import apply_pptx_structure_plan


@pytest.fixture
def deck_spec():
    path = Path(__file__).resolve().parents[1] / "examples/pptx/operational-review.json"
    return json.loads(path.read_text(encoding="utf-8"))


def _hash_parts(path):
    with zipfile.ZipFile(path) as archive:
        return {
            name: hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if not name.endswith("/")
        }


def test_builds_structured_deck(tmp_path, deck_spec):
    output = build_pptx(deck_spec, tmp_path / "review.pptx")
    package = inspect_ooxml(output)
    assert package["summary"]["errors"] == 0, package["issues"]
    extracted = extract_pptx(output)
    assert extracted["totals"]["slides"] == 9
    assert extracted["totals"]["charts"] == 1
    assert extracted["totals"]["tables"] == 1
    assert extracted["totals"]["notes_slides"] == 9
    assert extracted["slide_width"] == pytest.approx(13.3333, abs=0.001)
    report = validate_pptx(output, expected_slides=9, require_titles=True, require_notes=True)
    assert report["passed"] is True, report["issues"]
    assert report["summary"] == {"errors": 0, "warnings": 0, "info": 0}


def test_direct_text_edit_preserves_unrelated_parts(tmp_path, deck_spec):
    source = build_pptx(deck_spec, tmp_path / "review.pptx")
    output = tmp_path / "updated.pptx"
    before = _hash_parts(source)
    result = apply_pptx_edit_plan(
        source,
        output,
        {
            "replacements": [
                {
                    "find": "две пилотные команды",
                    "replace": "три пилотные команды",
                    "expected": 1,
                    "slides": [1],
                    "include_notes": True,
                },
                {"find": "две команды", "replace": "три команды", "expected": 1, "slides": [9]},
            ],
            "metadata": {"comments": "Updated scenario"},
        },
    )
    assert sum(item["count"] for item in result["replacements"]) == 2
    extracted = extract_pptx(output)
    assert "три команды" in extracted["slides"][8]["text"]
    assert "три пилотные команды" in extracted["slides"][0]["notes"]
    after = _hash_parts(output)
    changed = {name for name in before.keys() & after.keys() if before[name] != after[name]}
    assert changed == {"docProps/core.xml", "ppt/slides/slide9.xml", "ppt/notesSlides/notesSlide1.xml"}
    report = validate_pptx(output, expected_slides=9, require_notes=True)
    assert report["passed"] is True, report["issues"]


def test_replaces_text_split_across_drawing_runs(tmp_path):
    source = tmp_path / "split.pptx"
    output = tmp_path / "filled.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    paragraph = box.text_frame.paragraphs[0]
    first = paragraph.add_run()
    first.text = "Client: {{CLI"
    first.font.bold = True
    second = paragraph.add_run()
    second.text = "ENT_NAME}}"
    second.font.size = Pt(20)
    deck.save(source)

    apply_pptx_edit_plan(
        source,
        output,
        {"replacements": [{"find": "{{CLIENT_NAME}}", "replace": "Northern Systems", "expected": 1}]},
    )
    extracted = extract_pptx(output)
    assert "Client: Northern Systems" in extracted["slides"][0]["text"]
    assert inspect_ooxml(output)["summary"]["errors"] == 0


def test_media_replacement_keeps_part_identity(tmp_path):
    original_image = tmp_path / "original.png"
    replacement_image = tmp_path / "replacement.png"
    Image.new("RGB", (800, 450), "#23465C").save(original_image)
    Image.new("RGB", (800, 450), "#C77932").save(replacement_image)
    spec = {
        "slides": [
            {
                "layout": "image",
                "title": "Image",
                "image": original_image.name,
                "notes": "Image replacement test",
            }
        ]
    }
    source = build_pptx(spec, tmp_path / "image.pptx", base_dir=tmp_path)
    with zipfile.ZipFile(source) as archive:
        media_part = next(name for name in archive.namelist() if name.startswith("ppt/media/"))
        expected = hashlib.sha256(archive.read(media_part)).hexdigest()
    output = tmp_path / "image-updated.pptx"
    result = apply_pptx_edit_plan(
        source,
        output,
        {
            "media_replacements": [
                {"part": media_part, "path": replacement_image.name, "expected_sha256": expected}
            ]
        },
        base_dir=tmp_path,
    )
    assert result["media_replacements"][0]["dimensions"] == [800, 450]
    assert inspect_ooxml(output)["summary"]["errors"] == 0


def test_structural_slide_operations_and_orphan_cleanup(tmp_path):
    source = build_pptx(
        {
            "slides": [
                {"layout": "title", "title": "Opening"},
                {
                    "layout": "chart",
                    "title": "Chart",
                    "chart_type": "column",
                    "categories": ["A", "B"],
                    "series": [{"name": "Values", "values": [1, 2]}],
                },
                {"layout": "bullets", "title": "Close", "bullets": ["Done"]},
            ]
        },
        tmp_path / "structural.pptx",
    )
    output = tmp_path / "structural-updated.pptx"
    result = apply_pptx_structure_plan(
        source,
        output,
        {
            "duplicate": [{"slide": 3, "after": 3}],
            "delete": [2],
            "order": [3, 1, 2],
            "clean_orphans": True,
        },
    )
    assert result["slide_count"] == 3
    assert any(part.startswith("ppt/charts/") for part in result["orphaned_parts_removed"])
    assert any(part.startswith("ppt/embeddings/") for part in result["orphaned_parts_removed"])
    package = inspect_ooxml(output)
    assert package["summary"]["errors"] == 0, package["issues"]
    deck = extract_pptx(output)
    assert [slide["title"] for slide in deck["slides"]] == ["Close", "Opening", "Close"]
    assert deck["masters"] and deck["masters"][0]["layouts"]


def test_expected_count_mismatch_is_atomic(tmp_path, deck_spec):
    source = build_pptx(deck_spec, tmp_path / "review.pptx")
    output = tmp_path / "not-created.pptx"
    with pytest.raises(PptxEditError, match="expected 2"):
        apply_pptx_edit_plan(
            source,
            output,
            {"replacements": [{"find": "Операционная эффективность", "replace": "X", "expected": 2}]},
        )
    assert not output.exists()
