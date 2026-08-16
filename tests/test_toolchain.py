#!/usr/bin/env python3
"""
Regression tests for the document toolchain.

Every test here corresponds to something that was actually broken during
development and would fail silently if it regressed. Run:

    .venv/bin/python tests/test_toolchain.py
    .venv/bin/python tests/test_toolchain.py -v      # show each check
    .venv/bin/python tests/test_toolchain.py --fast  # skip Chromium renders

No pytest dependency: this must run with nothing but the toolchain installed.
"""

from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
PY = REPO / ".venv" / "bin" / "python"
if not PY.exists():
    PY = Path(sys.executable)

PASS, FAIL, SKIP = "\033[32m✓\033[0m", "\033[31m✗\033[0m", "\033[33m–\033[0m"
results: list[tuple[str, bool, str]] = []
VERBOSE = False
FAST = False


def check(name: str):
    """Decorator registering a test function."""
    def wrap(fn):
        fn._test_name = name
        return fn
    return wrap


def run(cmd: list[str], cwd: Path | None = None) -> subprocess.CompletedProcess:
    return subprocess.run(
        [str(c) for c in cmd], cwd=str(cwd or REPO),
        capture_output=True, text=True,
    )


# --------------------------------------------------------------------- tests
@check("recalc computes formulas and writes cached values")
def test_recalc(tmp: Path) -> None:
    import openpyxl

    path = tmp / "calc.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Модель"          # non-ASCII sheet name: engine uppercases it
    ws["A1"] = 100
    ws["A2"] = 250
    ws["A3"] = "=SUM(A1:A2)"
    ws["A4"] = "=A3*2"
    ws["A5"] = '=IFERROR(A4/0,"н/д")'
    ws["A6"] = "=INDEX(A1:A2,MATCH(250,A1:A2,0))"
    wb.save(path)

    proc = run([PY, REPO / "tools" / "recalc.py", path])
    assert proc.returncode == 0, f"recalc failed: {proc.stdout}{proc.stderr}"

    values = openpyxl.load_workbook(path, data_only=True)["Модель"]
    assert values["A3"].value == 350, f"A3 = {values['A3'].value}, expected 350"
    assert values["A4"].value == 700, f"A4 = {values['A4'].value}, expected 700"
    assert values["A5"].value == "н/д", f"A5 = {values['A5'].value!r}"
    assert values["A6"].value == 250, f"A6 = {values['A6'].value}"

    # formulas must survive: a recalc that flattens them destroys the model
    formulas = openpyxl.load_workbook(path, data_only=False)["Модель"]
    assert formulas["A3"].value == "=SUM(A1:A2)", "recalc destroyed the formula"


@check("recalc reports Excel errors and exits non-zero")
def test_recalc_errors(tmp: Path) -> None:
    import openpyxl

    path = tmp / "err.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    wb.active["A2"] = "=A1/0"
    wb.save(path)

    proc = run([PY, REPO / "tools" / "recalc.py", path])
    assert proc.returncode == 1, "an Excel error must exit 1"
    assert "DIV/0" in proc.stdout, f"error not reported: {proc.stdout}"


@check("recalc names functions the engine cannot evaluate")
def test_recalc_unsupported(tmp: Path) -> None:
    import openpyxl

    path = tmp / "xl.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    wb.active["A2"] = "=XLOOKUP(1,A1:A1,A1:A1)"
    wb.save(path)

    proc = run([PY, REPO / "tools" / "recalc.py", path, "--json"])
    assert "XLOOKUP" in proc.stdout, "unsupported function not named"


@check("validate detects a dangling relationship")
def test_validate_dangling(tmp: Path) -> None:
    import docx

    good = tmp / "good.docx"
    document = docx.Document()
    document.add_heading("Заголовок", 1)
    document.add_paragraph("Текст документа.")
    document.save(good)

    assert run([PY, REPO / "tools" / "validate.py", good]).returncode == 0

    broken = tmp / "broken.docx"
    with zipfile.ZipFile(good) as zin:
        items = {n: zin.read(n) for n in zin.namelist()}
    del items["word/styles.xml"]          # still referenced by document.xml.rels
    with zipfile.ZipFile(broken, "w") as zout:
        for name, blob in items.items():
            zout.writestr(name, blob)

    proc = run([PY, REPO / "tools" / "validate.py", broken])
    assert proc.returncode == 1, "missing part must fail validation"
    assert "missing part" in proc.stdout


@check("validate flags placeholder text")
def test_validate_placeholders(tmp: Path) -> None:
    import docx

    path = tmp / "ph.docx"
    document = docx.Document()
    document.add_paragraph("Подготовлено для [Company Name].")
    document.add_paragraph("TODO: уточнить цифры")
    document.save(path)

    proc = run([PY, REPO / "tools" / "validate.py", path])
    assert "TODO" in proc.stdout and "Company Name" in proc.stdout
    assert run([PY, REPO / "tools" / "validate.py", path, "--strict"]).returncode == 1, \
        "--strict must fail on warnings"


@check("validate catches uncached formulas")
def test_validate_uncached(tmp: Path) -> None:
    import openpyxl

    path = tmp / "nocache.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = 5
    wb.active["A2"] = "=A1*2"
    wb.save(path)

    proc = run([PY, REPO / "tools" / "validate.py", path])
    assert "no cached value" in proc.stdout, "uncached formulas not reported"


@check("validate reports shapes that overflow the slide")
def test_validate_overflow(tmp: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    path = tmp / "over.pptx"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(8), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "за краем слайда"
    prs.save(path)

    proc = run([PY, REPO / "tools" / "validate.py", path])
    assert "overflow" in proc.stdout.lower(), f"overflow not reported: {proc.stdout}"


@check("docx_fix repairs the missing Normal style")
def test_docx_fix(tmp: Path) -> None:
    # The script lives in the repo root so `require('docx')` resolves;
    # node searches upward from the *script's* directory, not the cwd.
    script = REPO / f".test-gen-{tmp.name}.js"
    out = tmp / "styled.docx"
    script.write_text(
        "const {Document,Packer,Paragraph,HeadingLevel}=require('docx');\n"
        "const fs=require('fs');\n"
        "const d=new Document({sections:[{children:["
        "new Paragraph({text:'Заголовок',heading:HeadingLevel.HEADING_1}),"
        "new Paragraph('текст')]}]});\n"
        f"Packer.toBuffer(d).then(b=>fs.writeFileSync({str(out)!r},b));\n",
        encoding="utf-8",
    )
    try:
        proc = run(["node", script])
        assert proc.returncode == 0, f"generation failed: {proc.stderr}"
    finally:
        script.unlink(missing_ok=True)

    import pypandoc

    before = pypandoc.convert_file(str(out), "markdown")
    assert "# Заголовок" not in before, (
        "docx-js unexpectedly produced a resolvable heading — "
        "if the library fixed this, docx_fix.js can be retired"
    )

    assert run(["node", REPO / "tools" / "js" / "docx_fix.js", out]).returncode == 0
    after = pypandoc.convert_file(str(out), "markdown")
    assert "# Заголовок" in after, "docx_fix did not restore heading structure"

    with zipfile.ZipFile(out) as zf:
        styles = zf.read("word/styles.xml").decode()
    assert 'w:styleId="Normal"' in styles
    assert 'w:outlineLvl w:val="0"' in styles
    assert run([PY, REPO / "tools" / "validate.py", out]).returncode == 0, \
        "docx_fix produced an invalid package"


@check("chromium launches and renders Cyrillic")
def test_chromium(tmp: Path) -> None:
    if FAST:
        raise SkipTest("--fast")
    out = tmp / "render.pdf"
    script = (
        "const {renderPdf}=require('./tools/js/chromium.js');"
        f"renderPdf('<h1>Проверка</h1>', {str(out)!r})"
        ".then(()=>console.log('ok')).catch(e=>{console.error(e.message);process.exit(1)});"
    )
    # env -i style: prove it works without an inherited LD_LIBRARY_PATH
    proc = subprocess.run(
        ["node", "-e", script], cwd=str(REPO),
        capture_output=True, text=True,
        env={k: v for k, v in __import__("os").environ.items() if k != "LD_LIBRARY_PATH"},
    )
    assert proc.returncode == 0, f"chromium failed: {proc.stderr}"
    assert out.exists() and out.stat().st_size > 1000

    import pypdf

    text = pypdf.PdfReader(str(out)).pages[0].extract_text()
    assert "Проверка" in text, f"cyrillic not in PDF text layer: {text!r}"


@check("fonts have both Latin and Cyrillic coverage")
def test_fonts(tmp: Path) -> None:
    from fontTools.ttLib import TTFont

    font_dir = REPO / "assets" / "fonts"
    generated = sorted(
        f for f in font_dir.glob("*.ttf") if not f.name.startswith("DejaVu")
    )
    if not generated:
        raise SkipTest("fonts not installed — run tools/install_fonts.py")

    # ₽ lives in the latin-ext subset, not latin or cyrillic. Merging only
    # latin+cyrillic produced fonts that rendered every ruble as a blank box.
    required = {"A": 0x41, "А": 0x410, "₽": 0x20BD, "—": 0x2014,
                "«": 0x00AB, "№": 0x2116}
    for path in generated:
        font = TTFont(str(path), lazy=True, fontNumber=0)
        points: set[int] = set()
        for table in font["cmap"].tables:
            points |= set(table.cmap)
        missing = [ch for ch, cp in required.items() if cp not in points]
        assert not missing, f"{path.name}: missing {' '.join(missing)}"


@check("render produces one page per slide, not two")
def test_render_slide_count(tmp: Path) -> None:
    if FAST:
        raise SkipTest("--fast")
    from pptx import Presentation
    from pptx.util import Inches

    path = tmp / "notes.pptx"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    for i in range(3):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        box.text_frame.text = f"Слайд {i + 1}"
        slide.notes_slide.notes_text_frame.text = "Заметки докладчика"
    prs.save(path)

    out = tmp / "qa"
    proc = run([PY, REPO / "tools" / "render.py", path, "-o", out])
    assert proc.returncode == 0, proc.stderr
    pages = sorted(out.glob("page-*.png"))
    assert len(pages) == 3, f"expected 3 pages for 3 slides, got {len(pages)}"


@check("render keeps slide aspect ratio")
def test_render_aspect(tmp: Path) -> None:
    if FAST:
        raise SkipTest("--fast")
    from pptx import Presentation
    from pptx.util import Inches

    path = tmp / "ar.pptx"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "x"
    prs.save(path)

    out = tmp / "qa"
    assert run([PY, REPO / "tools" / "render.py", path, "-o", out, "--pdf-only"]).returncode == 0

    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(str(out / "ar.pdf"))
    try:
        width, height = doc[0].get_size()
    finally:
        doc.close()
    ratio = width / height
    assert abs(ratio - 16 / 9) < 0.02, f"aspect {ratio:.3f}, expected 1.778 (16:9)"


@check("render shows a Title-styled paragraph")
def test_render_title(tmp: Path) -> None:
    import docx
    import pypandoc

    path = tmp / "title.docx"
    document = docx.Document()
    document.add_paragraph("Название документа", style="Title")
    document.add_paragraph("Основной текст.")
    document.save(path)

    # A fragment conversion drops it; --standalone keeps it. render.py must
    # use the latter, or the most prominent line vanishes from the preview.
    full = pypandoc.convert_file(str(path), "html5", extra_args=["--standalone"])
    assert "Название документа" in full


@check("examples run and produce valid output")
def test_examples(tmp: Path) -> None:
    if FAST:
        raise SkipTest("--fast")
    work = REPO / ".workdir"

    assert run(["node", REPO / "examples" / "build_report.js"]).returncode == 0
    assert run([PY, REPO / "examples" / "build_model.py"]).returncode == 0
    assert run(["node", REPO / "examples" / "build_deck.js"]).returncode == 0

    for name in ("Квартальный_отчёт.docx", "Финансовая_модель.xlsx", "Итоги_квартала.pptx"):
        path = work / name
        assert path.is_file(), f"{name} not produced"
        proc = run([PY, REPO / "tools" / "validate.py", path])
        assert proc.returncode == 0, f"{name} failed validation:\n{proc.stdout}"


@check("recalc leaves engine-gap formulas for Excel instead of caching #NAME?")
def test_recalc_engine_gap(tmp: Path) -> None:
    """
    SUBTOTAL is valid Excel but unimplemented by `formulas`. Caching the
    engine's #NAME? into the file turns a working workbook into a broken one,
    so such cells must be left uncached for Excel to fill on open.
    """
    import openpyxl

    path = tmp / "subtotal.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Реестр"
    for i, v in enumerate([10, 20, 30], start=1):
        ws.cell(i, 1, v)
    ws["A4"] = "=SUBTOTAL(109,A1:A3)"   # engine gap
    ws["B1"] = "=A1*2"                  # ordinary formula, must be cached
    wb.save(path)

    proc = run([PY, REPO / "tools" / "recalc.py", path, "--json"])
    assert proc.returncode == 0, f"recalc should not fail on SUBTOTAL:\n{proc.stdout}"

    import json

    report = json.loads(proc.stdout)
    assert report["status"] == "success", report
    assert report["total_left_to_excel"] == 1, report
    assert "SUBTOTAL" in report["left_to_excel"][0], report

    with zipfile.ZipFile(path) as z:
        sheet = next(n for n in z.namelist() if n.startswith("xl/worksheets/sheet"))
        xml = z.read(sheet).decode()
    assert "#NAME?" not in xml, "recalc cached a fake error over a valid formula"
    assert 'fullCalcOnLoad="1"' in z_read_workbook(path), "Excel will not recalc on open"

    wb2 = openpyxl.load_workbook(path, data_only=True)
    assert wb2["Реестр"]["B1"].value == 20, "ordinary formula lost its cached value"


def z_read_workbook(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        return z.read("xl/workbook.xml").decode()


@check("showcase documents build and validate")
def test_showcase(tmp: Path) -> None:
    if FAST:
        raise SkipTest("--fast")
    proc = run([PY, REPO / "showcase" / "build_all.py"])
    assert proc.returncode == 0, f"showcase build failed:\n{proc.stdout}{proc.stderr}"

    out = REPO / "showcase" / "out"
    expected = [
        "Модернизация_сети_накопителей.docx",
        "Финансовая_модель_накопители.xlsx",
        "Совет_директоров_накопители.pptx",
        "Резюме_программы_одна_страница.pdf",
    ]
    for name in expected:
        assert (out / name).is_file(), f"{name} not produced"

    # The docx is the one that exercises footnotes, comments, tracked changes
    # and OMML — verify the parts survived rather than trusting the exit code.
    with zipfile.ZipFile(out / expected[0]) as z:
        names = z.namelist()
        doc = z.read("word/document.xml").decode()
    for part in ("word/footnotes.xml", "word/comments.xml", "word/numbering.xml"):
        assert part in names, f"missing {part}"
    for marker in ("<w:footnoteReference", "<w:commentRangeStart", "<w:ins ",
                   "<w:del ", "<m:oMath>", 'w:orient="landscape"', "<w:vMerge"):
        assert marker in doc, f"docx lost {marker}"


@check("setup --check passes")
def test_setup_check(tmp: Path) -> None:
    proc = run(["bash", REPO / "tools" / "setup.sh", "--check"])
    assert proc.returncode == 0, f"setup --check failed:\n{proc.stdout}{proc.stderr}"
    assert "chromium renders PDF" in proc.stdout


class SkipTest(Exception):
    pass


def main() -> int:
    global VERBOSE, FAST
    ap = argparse.ArgumentParser()
    ap.add_argument("-v", "--verbose", action="store_true")
    ap.add_argument("--fast", action="store_true", help="skip Chromium-dependent tests")
    args = ap.parse_args()
    VERBOSE, FAST = args.verbose, args.fast

    tests = [
        obj for name, obj in sorted(globals().items())
        if callable(obj) and hasattr(obj, "_test_name")
    ]

    failures = skipped = 0
    for fn in tests:
        tmp = Path(tempfile.mkdtemp(prefix="doctest-"))
        try:
            fn(tmp)
            print(f"{PASS} {fn._test_name}")
        except SkipTest as exc:
            skipped += 1
            print(f"{SKIP} {fn._test_name} ({exc})")
        except AssertionError as exc:
            failures += 1
            print(f"{FAIL} {fn._test_name}\n    {exc}")
        except Exception as exc:  # noqa: BLE001
            failures += 1
            print(f"{FAIL} {fn._test_name}\n    {type(exc).__name__}: {exc}")
            if VERBOSE:
                import traceback
                traceback.print_exc()
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    total = len(tests)
    print(
        f"\n{total - failures - skipped}/{total} passed"
        + (f", {skipped} skipped" if skipped else "")
        + (f", \033[31m{failures} failed\033[0m" if failures else "")
    )
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
