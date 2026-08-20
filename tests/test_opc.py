from __future__ import annotations

import zipfile

from docx import Document

from document_system.opc import inspect_ooxml


def _rewrite(source, output, *, drop=(), additions=None):
    drop = set(drop)
    additions = additions or {}
    with zipfile.ZipFile(source) as incoming, zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as outgoing:
        for info in incoming.infolist():
            if info.filename not in drop:
                outgoing.writestr(info, incoming.read(info.filename))
        for name, value in additions.items():
            outgoing.writestr(name, value)


def test_detects_broken_internal_relationship(tmp_path):
    source = tmp_path / "source.docx"
    broken = tmp_path / "broken.docx"
    document = Document()
    document.add_paragraph("Hello")
    document.save(source)
    _rewrite(source, broken, drop={"word/styles.xml"})
    report = inspect_ooxml(broken)
    codes = {item["code"] for item in report["issues"]}
    assert "broken-relationship" in codes
    assert "missing-overridden-part" in codes
    assert report["summary"]["errors"] >= 2


def test_inventories_macro_content(tmp_path):
    source = tmp_path / "source.docx"
    macro = tmp_path / "macro.docx"
    document = Document()
    document.add_paragraph("Hello")
    document.save(source)
    _rewrite(source, macro, additions={"word/vbaProject.bin": b"not executable in test"})
    report = inspect_ooxml(macro)
    assert "word/vbaProject.bin" in report["package"]["macro_parts"]
    assert any(item["code"] == "macro-present" for item in report["issues"])
